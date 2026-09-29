import json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

D = sys.argv[1] if len(sys.argv) > 1 else '.'
SUBJ = [('marketing', 'マーケティング'), ('ai', 'AI演習（実践）'), ('literacy', 'ビジネス情報リテラシー')]
thin = Side(style='thin', color='999999')
B = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill('solid', fgColor='D9E2F3')
WRAP = Alignment(wrap_text=True, vertical='top')
CEN = Alignment(horizontal='center', vertical='center', wrap_text=True)

wb = Workbook(); wb.remove(wb.active)
for key, title in SUBJ:
    d = json.load(open(f'{D}/{key}.json'))
    ws = wb.create_sheet(title)
    for col, w in zip('ABCDEFGH', [6, 18, 26, 34, 24, 18, 14, 22]):
        ws.column_dimensions[col].width = w

    def cell(r, c, v, fill=None, al=WRAP, bold=False):
        x = ws.cell(r, c, v); x.alignment = al; x.border = B
        if fill: x.fill = fill
        if bold: x.font = Font(bold=True)
        return x

    def block(r1, r2, c1, c2, v, fill=None, al=WRAP, bold=False):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        cell(r1, c1, v, fill, al, bold)
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                ws.cell(r, c).border = B

    block(1, 1, 1, 8, 'AIビジネス専門学校　シラバス', al=CEN, bold=True)
    ws.cell(1, 1).font = Font(bold=True, size=14)
    h = d['header']
    for r, pairs in [(3, [('年度', h['年度']), ('学科', h['学科']), ('学年', h['学年']), ('総授業時間', h['総授業時間'])]),
                     (4, [('科目名', h['科目名']), ('担当教員', h['担当教員']), ('授業形態', h['授業形態']), ('単位数', h['単位数'])])]:
        for i, (k, v) in enumerate(pairs):
            cell(r, 1 + 2 * i, k, HEAD, CEN, True); cell(r, 2 + 2 * i, v, al=CEN)
    for r0, (lk, lv), (rk, rv) in [(5, ('科目の概要・位置づけ', d['概要']), ('使用教材・必要な環境', d['教材'])),
                                   (11, ('到達目標', d['到達目標']), ('評価基準・単位認定基準', d['評価基準'])),
                                   (17, ('成績評価の方法・割合', d['評価割合']), ('課題へのフィードバック方法', d['フィードバック']))]:
        block(r0, r0, 1, 4, lk, HEAD, CEN, True); block(r0, r0, 5, 8, rk, HEAD, CEN, True)
        block(r0 + 1, r0 + 5, 1, 4, lv); block(r0 + 1, r0 + 5, 5, 8, rv)
        for r in range(r0 + 1, r0 + 6):
            ws.row_dimensions[r].height = 30
    block(23, 23, 1, 8, '授業計画', HEAD, CEN, True)
    cols = ['回', 'テーマ', '到達目標', '授業内容・方法', '課題・提出物', '理解確認', '評価', '備考']
    for i, k in enumerate(cols):
        cell(24, 1 + i, k, HEAD, CEN, True)
    for j, s in enumerate(d['sessions']):
        r = 25 + j
        for i, k in enumerate(cols):
            cell(r, 1 + i, s[k], al=CEN if k == '回' else WRAP)
        ws.row_dimensions[r].height = 90
    ws.freeze_panes = 'A25'
wb.save(f'{D}/syllabus.xlsx')
print('ok', [(k, len(json.load(open(f"{D}/{k}.json"))["sessions"])) for k, _ in SUBJ])
