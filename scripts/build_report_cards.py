"""通知表を全学生分作る（1人1ページの PDF。2026年度前期）。

    python3 scripts/build_report_cards.py [成績表のAI計算版.xlsx]
    python3 scripts/build_report_cards.py --certificate AIBC26056,AIBC26059 [成績表のAI計算版.xlsx]   # 指定した人の成績証明書（塗らない）
    python3 scripts/build_report_cards.py --sample [成績表のAI計算版.xlsx]   # 合否がまばらな3人で見本4種類

入力を省くと data/output/成績表_2026前期_AI計算版.xlsx（build_ai_version.py の出力）を読む。
点数・評定・氏名は AI計算版シートの値。Excel 等で保存した計算済みの値があればそれを、無ければ LibreOffice で再計算して読む。
出席は同じブックに写した出席簿から数える。
出力は data/output/通知表_2026前期.pdf と、同じ内容の .html。発行日は実行した日。
決まりは src/grading/export/report_card.py の冒頭を参照。分からないものがあれば作らずに止まる。
"""

import datetime as dt
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "scripts"))

from openpyxl import load_workbook  # noqa: E402

from build_ai_version import DATE_CORRECTIONS, OUTPUT, recalculated  # noqa: E402
from grading.export.report_card import ReportCardError, read_cards, render_html  # noqa: E402
from grading.importing.attendance import read_register, subject_counts  # noqa: E402

DEPTS = [  # （AI計算版のシート, 学科名, 出席簿シートの頭）
    ("AI計算版_国際", "国際ビジネス科", "国際ビジネスAI科"),
    ("AI計算版_総合", "総合ビジネス科", "総合ビジネス科"),
]
PDF = OUTPUT.with_name("通知表_2026前期.pdf")
CHROMIUM = ["chromium", "chromium-browser", "google-chrome", "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"]


def chromium() -> str:
    for name in [os.environ.get("CHROMIUM", ""), *CHROMIUM]:
        if name and (shutil.which(name) or Path(name).is_file()):
            return shutil.which(name) or name
    raise RuntimeError("Chromium が無いため PDF にできない（環境変数 CHROMIUM で場所を指定できる）")


def computed(source: Path):
    """計算済みの値で読んだブック。式の値が保存されていない（openpyxl で書いたまま）なら再計算する。"""
    values = load_workbook(source, data_only=True)
    formulas = load_workbook(source)
    for sheet, _, _ in DEPTS:
        for row in formulas[sheet].iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("=") and values[sheet][cell.coordinate].value is None:
                    return recalculated(source)
    return values


SAMPLES = [  # （ファイル名, E を黒塗り, スタンプ, 位置）
    ("通知表_見本6_不合格_後期欄.pdf", True, "不合格", "body"),
    ("通知表_見本7_不合格_題名横.pdf", True, "不合格", "title"),
    ("通知表_見本8_あなたは不合格です_後期欄.pdf", True, "あなたは不合格です", "body"),
    ("通知表_見本9_あなたは不合格です_題名横.pdf", True, "あなたは不合格です", "title"),
]


def all_cards(source: Path):
    values = computed(source)
    cards = []
    for sheet, dept, prefix in DEPTS:
        reg = read_register(source, date_corrections=DATE_CORRECTIONS, sheet_prefix=prefix)
        if not reg.students:
            raise ReportCardError(f"{source.name} に「{prefix}」の出席簿シートが無い")
        cards += read_cards(values[sheet], dept, subject_counts(reg))
    return cards


def to_pdf(page_html: str, pdf: Path) -> Path:
    pdf.parent.mkdir(parents=True, exist_ok=True)
    page = pdf.with_suffix(".html")
    page.write_text(page_html, encoding="utf-8")
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([chromium(), "--headless", "--no-sandbox", "--disable-gpu", f"--user-data-dir={tmp}",
                        "--no-pdf-header-footer", f"--print-to-pdf={pdf}", page.as_uri()],
                       check=True, capture_output=True, timeout=300)
    return pdf


def build(source: Path, issued: dt.date) -> Path:
    """本番の形（利用者が見本⑦を選択、2026-10-02）: E の科目は黒地に白文字。E のある学生は題名の右に
    黒地に白文字の「不合格」スタンプをまっすぐ押す。"""
    cards = all_cards(source)
    print(f"{len(cards)}人分（うち不合格のスタンプ {sum(any(l.grade == 'E' for l in c.lines) for c in cards)}人）")
    return to_pdf(render_html(cards, issued, black_e=True, stamp="不合格", place="title", solid=True), PDF)


def certificates(source: Path, issued: dt.date, ids: list[str]) -> Path:
    """指定した学生だけの成績証明書（中身は通知表と同じ。黒塗り・スタンプ無し）。見つからない学籍番号があれば止める。"""
    cards = {c.student_id: c for c in all_cards(source)}
    missing = [i for i in ids if i not in cards]
    if missing:
        raise ReportCardError(f"成績表に無い学籍番号 {missing}")
    for i in ids:
        print(cards[i].dept, i, cards[i].name_ja)
    return to_pdf(render_html([cards[i] for i in ids], issued, kind="成績証明書"),
                  PDF.with_name(f"成績証明書_2026前期_{'_'.join(ids)}.pdf"))


def samples(source: Path, issued: dt.date, n: int = 3) -> list[Path]:
    """合否がまばらな（E の科目数が受講科目の半分に近い）n 人を、学科が偏らないよう交互に選び、見本を作る。"""
    def spread(c):
        return abs(sum(l.grade == "E" for l in c.lines) - len(c.lines) / 2), c.student_id
    by_dept = [sorted((c for c in all_cards(source) if c.dept == dept), key=spread) for _, dept, _ in DEPTS]
    worst = [by_dept[i % len(by_dept)][i // len(by_dept)] for i in range(n)]
    for c in worst:
        print(c.dept, c.student_id, "E", sum(l.grade == "E" for l in c.lines), "科目")
    return [to_pdf(render_html(worst, issued, black_e, stamp, place), PDF.with_name(name))
            for name, black_e, stamp, place in SAMPLES]


if __name__ == "__main__":
    argv = sys.argv[1:]
    ids = []
    if "--certificate" in argv:
        k = argv.index("--certificate")
        ids = argv[k + 1].split(",")
        del argv[k:k + 2]
    args = [a for a in argv if a != "--sample"]
    src = Path(args[0]) if args else OUTPUT
    if ids:
        print(certificates(src, dt.date.today(), ids))
    elif "--sample" in sys.argv:
        print(*samples(src, dt.date.today()), sep="\n")
    else:
        print(build(src, dt.date.today()))
