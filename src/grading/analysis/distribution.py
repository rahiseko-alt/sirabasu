"""度数分布（県の学力検査の報告の形にならう）。総点は50点刻み、科目の合計は10点刻み。折れ線グラフつき。"""

import statistics
from collections.abc import Sequence

from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill

from grading.export.e_list import GradeRow

_HEAD = PatternFill("solid", fgColor="DDE7F3")


def _bins(values, width, top):
    edges = list(range(0, top, width))
    counts = [0] * len(edges)
    for v in values:
        counts[min(int(v // width), len(edges) - 1)] += 1
    labels = [f"{e}〜{e + width - 1}" if e + width < top else f"{e}〜{top}" for e in edges]
    return labels, counts


def student_totals(rows: Sequence[GradeRow]) -> dict[str, float]:
    out: dict[str, float] = {}
    for r in rows:
        if isinstance(r.total, (int, float)):
            out[r.student_id] = out.get(r.student_id, 0) + r.total
    return out


def _table(ws, start_row, title, labels, columns, chart_title, x_title, anchor):
    ws.cell(start_row, 1, title).font = Font(bold=True, size=12)
    head = start_row + 1
    ws.cell(head, 1, "段階点")
    for k, (name, _) in enumerate(columns):
        ws.cell(head, 2 + k, name)
    for c in range(1, 2 + len(columns)):
        ws.cell(head, c).font, ws.cell(head, c).fill = Font(bold=True), _HEAD
    for i, label in enumerate(labels):
        ws.cell(head + 1 + i, 1, label)
        for k, (_, counts) in enumerate(columns):
            ws.cell(head + 1 + i, 2 + k, counts[i])
    chart = LineChart()
    chart.title, chart.y_axis.title, chart.x_axis.title = chart_title, "人数", x_title
    chart.height, chart.width = 7.5, 16
    data = Reference(ws, min_col=2, max_col=1 + len(columns), min_row=head, max_row=head + len(labels))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=head + 1, max_row=head + len(labels)))
    ws.add_chart(chart, anchor)
    return head + len(labels) + 2


def add_distribution_sheet(wb, groups: Sequence[tuple[str, Sequence[GradeRow]]], title: str = "度数分布"):
    ws = wb.create_sheet(title)
    totals = {dept: student_totals(rows) for dept, rows in groups}
    n_subjects = max(len({r.subject for r in rows if r.student_id == sid}) for _, rows in groups
                     for sid in student_totals(rows))
    top = 100 * n_subjects
    cols = [(dept, _bins(t.values(), 50, top)[1]) for dept, t in totals.items()]
    allv = [v for t in totals.values() for v in t.values()]
    labels, both = _bins(allv, 50, top)
    cols.append(("両学科", both))
    row = _table(ws, 1, f"総点の度数分布（{n_subjects}科目・{top}点満点）", labels, cols, "総点の分布", "総点", "F2")
    ws.cell(row, 1, "総点の統計").font = Font(bold=True)
    ws.append(["", "人数", "平均", "標準偏差", "最高", "最低"])
    for name, values in [*((d, list(t.values())) for d, t in totals.items()), ("両学科", allv)]:
        ws.append([name, len(values), statistics.fmean(values), statistics.pstdev(values), max(values), min(values)])
        for c in (3, 4):
            ws.cell(ws.max_row, c).number_format = "0.0"
    row = ws.max_row + 2
    for dept, rows in groups:
        subjects = list(dict.fromkeys(r.subject for r in rows))
        cols = [(s, _bins([r.total for r in rows if r.subject == s and isinstance(r.total, (int, float))], 10, 100)[1])
                for s in subjects]
        labels = _bins([], 10, 100)[0]
        row = max(_table(ws, row, f"{dept} 科目別の合計点の度数分布（100点満点）", labels, cols,
                         f"{dept} 科目別の分布", "合計点", f"F{row + 1}"), row + 18)
    ws.column_dimensions["A"].width = 14
    for col in "BCDE":
        ws.column_dimensions[col].width = 14
    return ws
