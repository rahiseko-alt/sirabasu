"""成績分析のシート（テスト分析・偏差値）を作る。

成績表（値）の「〜テスト」という評価項目を、5行目の点数配分を満点として分析する。空欄の学生（未受験・未履修）は除く。
"""

import statistics
from collections.abc import Sequence
from dataclasses import dataclass

from openpyxl.styles import Alignment, Font, PatternFill

from grading.analysis.test_stats import correlation, deviation_scores, judge, summarize
from grading.export.fill import _columns

_HEAD = PatternFill("solid", fgColor="DDE7F3")
_VERDICT = {"適正": "E2EFDA", "易しすぎ": "FFF2CC", "難しすぎ": "F8CBAD", "差がつかない": "FCE4D6"}


@dataclass(frozen=True)
class TestColumn:
    dept: str
    subject: str
    item: str
    maximum: float
    scores: dict[str, float]          # 学籍番号 → 得点
    names: dict[str, tuple[str, str]]


def collect_tests(ws, dept: str) -> list[TestColumn]:
    out = []
    names = {str(ws.cell(r, 2).value).strip(): (ws.cell(r, 3).value, " ".join(str(ws.cell(r, 4).value or "").split()))
             for r in range(6, ws.max_row + 1) if ws.cell(r, 2).value}
    for (subject, item), col in _columns(ws).items():
        if "テスト" not in item:
            continue
        maximum = ws.cell(5, col).value
        if not isinstance(maximum, (int, float)) or maximum <= 0:
            continue
        scores = {}
        for r in range(6, ws.max_row + 1):
            sid, v = ws.cell(r, 2).value, ws.cell(r, col).value
            if sid and isinstance(v, (int, float)):
                scores[str(sid).strip()] = float(v)
        if scores:
            out.append(TestColumn(dept, ws.cell(3, _subject_col(ws, col)).value.strip(), item.replace("\n", ""),
                                  float(maximum), scores, names))
    return out


def _subject_col(ws, col):
    while col > 1 and not (isinstance(ws.cell(3, col).value, str) and ws.cell(3, col).value.strip()):
        col -= 1
    return col


def _other_rate(tests: Sequence[TestColumn], t: TestColumn, sid: str):
    rates = [o.scores[sid] / o.maximum for o in tests if o.subject != t.subject and o.dept == t.dept and sid in o.scores]
    return statistics.fmean(rates) if rates else None


def add_test_analysis(wb, tests: Sequence[TestColumn], title: str = "テスト分析"):
    ws = wb.create_sheet(title)
    ws.append(["テストの適正度（設問別の得点が無いため、テスト全体の得点で判定）"])
    ws.append(["判定の目安: 易しすぎ＝平均得点率85%以上か天井効果（平均＋標準偏差＞満点）／難しすぎ＝平均得点率40%未満か床効果（平均−標準偏差＜0）"
               "／差がつかない＝標準偏差が満点の10%未満。他テストとの相関が低い（0.3未満）テストは、ほかの力と違うものを測っている可能性がある"])
    head = ["学科", "科目", "テスト", "判定", "理由", "受験者数", "満点", "平均", "平均得点率", "標準偏差", "最高", "最低", "中央値",
            "満点者数", "0点者数", "他テストとの相関", *[f"{i * 10}〜{i * 10 + 10}%" for i in range(10)]]
    ws.append(head)
    for cell in ws[3]:
        cell.font, cell.fill, cell.alignment = Font(bold=True), _HEAD, Alignment(wrap_text=True, horizontal="center")
    groups = list(tests)
    for subject_item in dict.fromkeys((t.subject, t.item) for t in tests):
        same = [t for t in tests if (t.subject, t.item) == subject_item]
        if len(same) > 1 and len({t.maximum for t in same}) == 1:
            merged = {f"{t.dept}:{k}": v for t in same for k, v in t.scores.items()}
            groups.append(TestColumn("両学科", same[0].subject, same[0].item, same[0].maximum, merged, {}))
    groups.sort(key=lambda t: (t.subject, t.item, t.dept != "両学科", t.dept))
    for t in groups:
        s = summarize(list(t.scores.values()), t.maximum)
        verdict, reasons = judge(s)
        if t.dept == "両学科":
            r = None
        else:
            pairs = [(v / t.maximum, _other_rate(tests, t, sid)) for sid, v in t.scores.items()]
            pairs = [(a, b) for a, b in pairs if b is not None]
            r = correlation([a for a, _ in pairs], [b for _, b in pairs]) if pairs else None
        ws.append([t.dept, t.subject, t.item, verdict, "／".join(reasons), s.n, s.maximum, s.mean, s.rate, s.sd,
                   s.maximum_score, s.minimum, s.median, s.full_marks, s.zeros, r, *s.bins])
        row = ws.max_row
        ws.cell(row, 4).fill = PatternFill("solid", fgColor=_VERDICT[verdict])
        for c, fmt in ((8, "0.0"), (9, "0%"), (10, "0.0"), (13, "0.0"), (16, "0.00")):
            ws.cell(row, c).number_format = fmt
    for col, width in zip("ABCDE", (12, 26, 16, 12, 50)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "D4"
    return ws


def add_deviation_sheet(wb, tests: Sequence[TestColumn], title: str = "偏差値"):
    """学生ごとの各テストの偏差値（学科ごとに計算）。"""
    ws = wb.create_sheet(title)
    cols = [(t.dept, t.subject, t.item) for t in tests]
    ws.append(["学科", "学籍番号", "氏名", "カタカナ", *[f"{s}\n{i}" for _, s, i in cols], "平均偏差値"])
    for cell in ws[1]:
        cell.font, cell.fill, cell.alignment = Font(bold=True), _HEAD, Alignment(wrap_text=True, horizontal="center")
    devs = {(t.dept, t.subject, t.item): deviation_scores(t.scores) for t in tests}
    for dept in dict.fromkeys(t.dept for t in tests):
        names = next(t.names for t in tests if t.dept == dept)
        for sid, (name, kana) in names.items():
            values = [devs[c].get(sid) if c[0] == dept else None for c in cols]
            got = [v for v in values if v is not None]
            ws.append([dept, sid, name, kana, *values, statistics.fmean(got) if got else None])
            for c in range(5, 6 + len(cols)):
                ws.cell(ws.max_row, c).number_format = "0.0"
    for col, width in zip("ABCD", (12, 12, 30, 24)):
        ws.column_dimensions[col].width = width
    ws.row_dimensions[1].height = 48
    ws.freeze_panes = "E2"
    return ws
