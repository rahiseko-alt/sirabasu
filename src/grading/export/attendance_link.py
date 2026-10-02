"""出席簿を成績ファイルに丸ごと写し、出席点・出席率・態度点を出席簿のセルを数える式にする。

たどり方: AI計算版の出席点 → 出席集計シートの「出席点」→ 同じ行の 授業数・欠席・遅刻 → 出席簿の月別シートの 欠・遅 のセル。
- 授業数・欠席・遅刻は、出席簿の4行目（科目名）が一致する列の「×以外」「欠」「遅」を数える式（SUMPRODUCT）
- 日本語運用力強化演習は「週ごと」シートで週ごとに 出・遅・欠 を式で判定してから数える（その週に1回でも出席なら出）
- 出席率＝1−（欠席＋遅刻÷3）÷授業数（出席簿にある式）。出席点は 60%未満で0点、4%下がるごとに減点
  （小数の誤差で段階がずれないよう、INT(100×(3×欠席＋遅刻)÷(12×授業数)) と整数で数える）
"""

import datetime as dt
import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

from grading.export.copy_sheet import copy_sheet
from grading.export.fill import AI_FILL, _columns, _subject
from grading.importing.attendance import Register

_HEAD = PatternFill("solid", fgColor="DDE7F3")
ITEMS = ["授業数", "欠席", "遅刻", "出席率", "出席点"]


def month_sheets(wb) -> list:
    return [ws for ws in wb.worksheets if re.search(r"\d+月$", ws.title.strip()) and "原紙" not in ws.title]


def copy_registers(register_wb, dst_wb) -> list[str]:
    """出席簿の月別シートを、シート名もそのままに写す。"""
    return [copy_sheet(ws, dst_wb, ws.title).title for ws in month_sheets(register_wb)]


def _q(title: str) -> str:
    return "'" + title.replace("'", "''") + "'"


def _text(value: str) -> str:
    return '"' + value.replace('"', '""') + '"'


@dataclass(frozen=True)
class Tally:
    sheet: str
    row: dict[str, int]                        # 学籍番号 → 行
    col: dict[tuple[str, str], int]            # (科目, 項目) → 列

    def ref(self, sid: str, subject: str, item: str) -> str:
        return f"{_q(self.sheet)}!${L(self.col[(subject, item)])}${self.row[sid]}"


def _student_rows(ws) -> dict[str, int]:
    return {str(ws.cell(r, 2).value).strip(): r for r in range(6, ws.max_row + 1) if ws.cell(r, 2).value}


def _end(ws) -> int:
    return next(c.column for c in ws[3] if c.value == "授業合計数")


def _count_formula(reg: Register, wb, subject: str, sid: str, condition: str) -> str:
    """科目名が一致する列のうち、condition（例 ="欠"）を満たすセルの数を、月ごとの SUMPRODUCT の和で表す。"""
    terms = []
    for ws in month_sheets(wb):
        heads = sorted({s.subject_original for s in reg.sessions if s.sheet == ws.title and s.subject == subject})
        rows = _student_rows(ws)
        if not heads or sid not in rows:
            continue
        last = L(_end(ws) - 1)
        head_rng = f"{_q(ws.title)}!$E$4:${last}$4"
        match = "+".join(f"({head_rng}={_text(h)})" for h in heads)
        match = f"({match})" if len(heads) == 1 else f"(({match})>0)"
        terms.append(f"SUMPRODUCT({match}*({_q(ws.title)}!$E${rows[sid]}:${last}${rows[sid]}{condition}))")
    return "=" + ("+".join(terms) if terms else "0")


def weekly_sheet(wb, reg: Register, subject: str, students: Sequence[tuple[str, str]], title: str) -> str:
    """週ごとの判定（出・遅・欠）を式で並べる。×だけの週は空欄。"""
    ws = wb.create_sheet(title)
    weeks: dict[dt.date, list] = {}
    for s in reg.sessions:
        if s.subject == subject and s.date is not None:
            weeks.setdefault(s.date - dt.timedelta(days=s.date.weekday()), []).append(s)
    ws.cell(1, 1, f"{subject}：週ごとの出欠（その週に1回でも出席なら「出」、出席が無く遅刻があれば「遅」、ほかは「欠」。×だけの週は空欄）")
    ws.cell(1, 1).font = Font(bold=True)
    ws.cell(2, 1, "学籍番号")
    ws.cell(2, 2, "氏名")
    mondays = sorted(weeks)
    for k, monday in enumerate(mondays):
        ws.cell(2, 3 + k, f"{monday:%m/%d}週")
    for c in ws[2]:
        c.font, c.fill = Font(bold=True), _HEAD
    sheets = {w.title: _student_rows(w) for w in month_sheets(wb)}
    for i, (sid, name_ref) in enumerate(students):
        r = 3 + i
        ws.cell(r, 1, sid)
        ws.cell(r, 2, f"={name_ref}")
        for k, monday in enumerate(mondays):
            cells = [f"{_q(s.sheet)}!{s.column}{sheets[s.sheet][sid]}" for s in weeks[monday] if sid in sheets[s.sheet]]
            if not cells:
                continue
            n = "+".join(f"({c}<>\"×\")" for c in cells)
            p = "+".join(f"({c}=\"\")" for c in cells)
            late = "+".join(f"({c}=\"遅\")" for c in cells)
            ws.cell(r, 3 + k, f'=IF(({n})=0,"",IF(({p})>0,"出",IF(({late})>0,"遅","欠")))')
            ws.cell(r, 3 + k).alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 12, 30
    ws.freeze_panes = "C3"
    return title


def tally_sheet(wb, reg: Register, rules: Mapping[str, tuple[int, int]], weekly: str,
                students: Sequence[tuple[str, str]], title: str, weekly_title: str) -> Tally:
    """学生×科目の 授業数・欠席・遅刻・出席率・出席点 を、出席簿を数える式で並べる。"""
    ws = wb.create_sheet(title)
    ws.cell(1, 1, "出席集計（出席簿の月別シートを数える式。出席率＝1−（欠席＋遅刻÷3）÷授業数、60%未満は0点、4%下がるごとに減点）")
    ws.cell(1, 1).font = Font(bold=True)
    ws.cell(3, 1, "学籍番号")
    ws.cell(3, 2, "氏名")
    col, c = {}, 3
    for subject in rules:
        ws.cell(2, c, subject + ("（1週＝1回）" if subject == weekly else ""))
        ws.merge_cells(start_row=2, start_column=c, end_row=2, end_column=c + len(ITEMS) - 1)
        for k, item in enumerate(ITEMS):
            ws.cell(3, c + k, item)
            col[(subject, item)] = c + k
        c += len(ITEMS)
    for row in (ws[2], ws[3]):
        for cell in row:
            cell.font, cell.fill = Font(bold=True), _HEAD
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
    rows = {}
    week_rows = {sid: 3 + i for i, (sid, _) in enumerate(students)}
    for i, (sid, name_ref) in enumerate(students):
        r = 4 + i
        rows[sid] = r
        ws.cell(r, 1, sid)
        ws.cell(r, 2, f"={name_ref}")
        for subject, (maximum, step) in rules.items():
            n, a, late, rate, pts = (L(col[(subject, item)]) for item in ITEMS)
            if subject == weekly:
                rng = f"{_q(weekly_title)}!$C${week_rows[sid]}:${L(wb[weekly_title].max_column)}${week_rows[sid]}"
                ws[f"{n}{r}"] = f'=COUNTIF({rng},"出")+COUNTIF({rng},"遅")+COUNTIF({rng},"欠")'
                ws[f"{a}{r}"] = f'=COUNTIF({rng},"欠")'
                ws[f"{late}{r}"] = f'=COUNTIF({rng},"遅")'
            else:
                ws[f"{n}{r}"] = _count_formula(reg, wb, subject, sid, '<>"×"')
                ws[f"{a}{r}"] = _count_formula(reg, wb, subject, sid, '="欠"')
                ws[f"{late}{r}"] = _count_formula(reg, wb, subject, sid, '="遅"')
            ws[f"{rate}{r}"] = f"=1-({a}{r}+{late}{r}/3)/{n}{r}"
            ws[f"{rate}{r}"].number_format = "0.0000%"
            ws[f"{pts}{r}"] = (f"=IF(5*(3*{n}{r}-3*{a}{r}-{late}{r})<3*3*{n}{r},0,"
                               f"{maximum}-{step}*INT(100*(3*{a}{r}+{late}{r})/(12*{n}{r})))")
            ws[f"{pts}{r}"].font = Font(bold=True)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 12, 30
    ws.row_dimensions[3].height = 30
    ws.freeze_panes = "C4"
    return Tally(title, rows, col)


def link_ai_cells(ai, tally: Tally, rate_cols: Mapping[str, int], rules: Mapping[str, tuple[int, int]],
                  late_deduct: Sequence[str], flat_ten: Sequence[str]) -> int:
    """AI計算版の出席点・出席率・態度点のセルを、出席集計を参照する式にして色を付ける。"""
    columns = _columns(ai)
    rows = _student_rows(ai)
    n = 0

    def put(r, c, value):
        nonlocal n
        ai.cell(r, c).value = value
        ai.cell(r, c).fill = AI_FILL
        n += 1

    for sid, r in rows.items():
        for subject in rules:
            put(r, columns[(_subject(subject), "出席")], f"={tally.ref(sid, subject, '出席点')}")
            if _subject(subject) in rate_cols:
                put(r, rate_cols[_subject(subject)], f"={tally.ref(sid, subject, '出席率')}")
        for subject in late_deduct:
            put(r, columns[(_subject(subject), "授業態度")], f"=10-2*{tally.ref(sid, subject, '遅刻')}")
        for subject in flat_ten:
            put(r, columns[(_subject(subject), "授業態度")], 10)
    return n
