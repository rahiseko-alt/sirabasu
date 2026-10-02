"""出席簿「科目ごと」シートの正しい集計（AI計算版）を書き出す。

元の「科目ごと」シートの並び（3行目に科目名、各科目5列: 授業合計数・欠席合計数・遅刻合計数・欠席数・出席率）を保つ。
- 授業合計数・欠席合計数・遅刻合計数: 月別シートからシステムが数え直した整数（根拠のセルは「根拠」シート）
- 欠席数・出席率: 元の出席簿と同じ式（欠席数＝欠席合計数＋遅刻合計数÷3、出席率＝1−欠席数÷授業合計数）を Excel の式で置く
元の値（再計算済み）との比較を「元の集計との比較」シートに並べる。
"""

from collections.abc import Mapping, Sequence
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from grading.importing.attendance import CalendarAnomaly, Problem, SubjectCount, unify_subject

AI_MARK = "（AI計算版）"
_DIFF = PatternFill("solid", fgColor="F8CBAD")


def _summary_sheet(wb):
    return next(ws for ws in wb.worksheets if "科目" in ws.title)


def write_subject_summary(
    src: str | Path,
    dst: str | Path,
    counts: Mapping[tuple[str, str], SubjectCount],
    original: Mapping[tuple[str, str], tuple] | None = None,
    problems: Sequence[Problem] = (),
    anomalies: Sequence[CalendarAnomaly] = (),
) -> Path:
    """src の「科目ごと」シートだけを残し、値を正しい集計に置き換えて dst に保存する。original は元の値（比較用）。"""
    wb = load_workbook(src)
    ws = _summary_sheet(wb)
    for other in [w for w in wb.worksheets if w is not ws]:
        wb.remove(other)
    ws.title = (ws.title + AI_MARK)[:31]
    if isinstance(ws["A1"].value, str):
        ws["A1"].value += AI_MARK

    groups = [(c, unify_subject(ws.cell(3, c).value)) for c in range(5, ws.max_column + 1)
              if isinstance(ws.cell(3, c).value, str) and ws.cell(3, c).value.strip()]
    cmp = wb.create_sheet("元の集計との比較")
    cmp.append(["学籍番号", "科目", "元の授業数", "正しい授業数", "元の欠席", "正しい欠席", "元の遅刻", "正しい遅刻", "差"])
    evidence = wb.create_sheet("根拠")
    evidence.append(["学籍番号", "科目", "区分", "セル（シート!セル）"])
    missing = []
    for r in range(6, ws.max_row + 1):
        sid = ws.cell(r, 2).value
        if not sid:
            continue
        sid = str(sid).strip()
        for c, subject in groups:
            n = counts.get((sid, subject))
            if n is None:
                missing.append((sid, subject))
                for k in range(5):
                    ws.cell(r, c + k).value = None
                continue
            ws.cell(r, c).value = n.sessions
            ws.cell(r, c + 1).value = n.absent
            ws.cell(r, c + 2).value = n.late
            a, b, d = (get_column_letter(c + k) for k in (0, 1, 2))
            h = get_column_letter(c + 3)
            ws.cell(r, c + 3).value = f"={b}{r}+{d}{r}/3"
            ws.cell(r, c + 4).value = f"=1-{h}{r}/{a}{r}"
            if original is not None:
                old = tuple(original.get((sid, subject), (None, None, None)))
                new = (n.sessions, n.absent, n.late)
                differs = old != new
                cmp.append([sid, subject, old[0], new[0], old[1], new[1], old[2], new[2], "違う" if differs else ""])
                if differs:
                    for k in range(3):
                        ws.cell(r, c + k).fill = _DIFF
            for cell in n.absent_cells:
                evidence.append([sid, subject, "欠", cell])
            for cell in n.late_cells:
                evidence.append([sid, subject, "遅", cell])

    check = wb.create_sheet("日程と記入の確認")
    check.append(["種類", "場所・日付", "内容"])
    for p in problems:
        check.append(["出席簿の記入", p.cell, p.detail])
    for a in anomalies:
        check.append([a.kind, f"{a.date:%Y-%m-%d}", a.detail])
    for sid, subject in missing:
        check.append(["集計できない", sid, f"{subject} の授業が月別シートに1回も無い（全て×）"])
    for sheet in (cmp, evidence, check):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        for col, width in zip("ABCDEFGHI", (12, 26, 14, 40, 10, 10, 10, 10, 6)):
            sheet.column_dimensions[col].width = width
    dst = Path(dst)
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    return dst
