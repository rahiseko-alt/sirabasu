"""科目ごとの評定E一覧。再計算済みの成績表（値）から読む。

その学生の評価項目がすべて空欄の科目（受講していない科目）は一覧に入れない。
"""

from collections.abc import Sequence
from dataclasses import dataclass

from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


@dataclass(frozen=True)
class GradeRow:
    subject: str
    student_id: str
    name: str
    total: object
    grade: object
    kana: str = ""


def _kana(value) -> str:
    return " ".join(str(value or "").replace("\n", " ").replace("\u3000", " ").split())


def grade_rows(ws) -> list[GradeRow]:
    groups, current = [], None
    for c in range(5, ws.max_column + 1):
        head = ws.cell(3, c).value
        if isinstance(head, str) and head.strip():
            current = {"subject": head.strip(), "items": [], "total": None, "grade": None}
            groups.append(current)
        item = ws.cell(4, c).value
        if current is None or not isinstance(item, str):
            continue
        if item.strip() == "合計":
            current["total"] = c
        elif item.strip() == "評定":
            current["grade"] = c
        elif current["total"] is None:
            current["items"].append(c)
    out = []
    for g in groups:
        if g["grade"] is None:
            continue
        for r in range(6, ws.max_row + 1):
            sid = ws.cell(r, 2).value
            if not sid:
                continue
            if all(ws.cell(r, c).value in (None, "") for c in g["items"]):
                continue
            out.append(GradeRow(g["subject"], str(sid).strip(), ws.cell(r, 3).value,
                                ws.cell(r, g["total"]).value if g["total"] else None, ws.cell(r, g["grade"]).value,
                                _kana(ws.cell(r, 4).value)))
    return out


_HEAD = PatternFill("solid", fgColor="DDE7F3")
_E = PatternFill("solid", fgColor="F8CBAD")


def _bold(row):
    for cell in row:
        cell.font, cell.fill = Font(bold=True), _HEAD


def _es(rows):
    return [r for r in rows if r.grade == "E"]


def add_e_matrix(wb, groups: Sequence[tuple[str, Sequence[GradeRow], dict[str, tuple[str, str]]]],
                 title: str = "E一覧"):
    """落ちこぼれ（評定E）の一覧。追試の予定向けに、Eのある学生×科目の表（セルはその科目の合計点）。

    groups: (学科, 評定の行, 科目→(担当, 曜日)) の並び。2学科以上なら「学科」の列を付け、
    学科で担当・曜日が違う科目は見出しに学科名を添える。見出しに担当・曜日・Eの人数。E科目数の多い順。
    """
    many = len(groups) > 1
    es = [(dept, e) for dept, rows, _ in groups for e in rows if e.grade == "E"]
    subjects = list(dict.fromkeys(e.subject for _, rows, _ in groups for e in rows if any(x.subject == e.subject for _, x in es)))

    def header(s, k):
        vals = {dept: info.get(s, ("", ""))[k] for dept, rows, info in groups if any(r.subject == s for r in rows)}
        if len(set(vals.values())) <= 1:
            return next(iter(vals.values()), "")
        return "／".join(f"{v}（{d[:2]}）" for d, v in vals.items())

    per: dict[tuple[str, str], list[GradeRow]] = {}
    for dept, e in es:
        per.setdefault((dept, e.student_id), []).append(e)
    order = sorted(per, key=lambda k: (-len(per[k]), k[0], k[1]))
    lead = ["学科"] if many else []
    pad = [""] * len(lead)
    ws = wb.create_sheet(title)
    ws.append([*pad, "科目", "", "", "", *subjects])
    ws.append([*pad, "担当", "", "", "", *[header(s, 0) for s in subjects]])
    ws.append([*pad, "曜日", "", "", "", *[header(s, 1) for s in subjects]])
    ws.append([*pad, "Eの人数", "", "", len(es), *[sum(1 for _, e in es if e.subject == s) for s in subjects]])
    ws.append([*lead, "学籍番号", "氏名", "カタカナ", "E科目数", *["合計点" for _ in subjects]])
    for r in range(1, 6):
        _bold(ws[r])
    first_subject_col = len(lead) + 5
    for key in order:
        items = per[key]
        by = {e.subject: e.total for e in items}
        ws.append([*([key[0]] if many else []), key[1], items[0].name, items[0].kana, len(items), *[by.get(s) for s in subjects]])
        for k, s in enumerate(subjects):
            if s in by:
                ws.cell(ws.max_row, first_subject_col + k).fill = _E
    for i, width in enumerate([*([14] if many else []), 12, 30, 24, 8]):
        ws.column_dimensions[get_column_letter(1 + i)].width = width
    ws.freeze_panes = f"{get_column_letter(first_subject_col)}6"
    return ws


def gpa_of(ws) -> dict[str, object]:
    """3行目が「GPA」の列から、学籍番号ごとのGPA（再計算済みの値）を読む。"""
    col = next((c for c in range(1, ws.max_column + 1) if str(ws.cell(3, c).value or "").strip() == "GPA"), None)
    if col is None:
        return {}
    return {str(ws.cell(r, 2).value).strip(): ws.cell(r, col).value for r in range(6, ws.max_row + 1) if ws.cell(r, 2).value}


def add_personal_grades(wb, rows: Sequence[GradeRow], gpa: dict[str, object], gpa_grade=None, title: str = "個人別評定"):
    """学生ごとに科目の評定とGPAだけを並べる。受講していない科目は「—」。gpa_grade は GPA→評定 の関数（未定なら空欄）。"""
    subjects = list(dict.fromkeys(r.subject for r in rows))
    students = list(dict.fromkeys((r.student_id, r.name, r.kana) for r in rows))
    grade = {(r.student_id, r.subject): r.grade for r in rows}
    ws = wb.create_sheet(title)
    ws.append(["学籍番号", "氏名", "カタカナ", *subjects, "GPA", "GPA評定"])
    for cell in ws[1]:
        cell.font, cell.fill = Font(bold=True), PatternFill("solid", fgColor="DDE7F3")
    for sid, name, kana in students:
        g = gpa.get(sid)
        ws.append([sid, name, kana, *[grade.get((sid, s), "—") for s in subjects], g,
                   gpa_grade(g) if gpa_grade and isinstance(g, (int, float)) else None])
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 24
    ws.freeze_panes = "D2"
    return ws
