"""白紙にした成績表（AI計算版の枠）に、システムが計算した値だけを入れる。

科目名（3行目）と評価項目名（4行目）と学籍番号（B列）からセルを探す。見つからなければ推測せずに止める。
入れた値はすべて「根拠」シートに、計算式と元のセルを添えて並べる。値は丸めない（割り切れない値は分数で残す）。
"""

import re
from collections.abc import Sequence
from dataclasses import dataclass
from fractions import Fraction
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

from grading.export.provenance import format_number
from grading.importing.attendance import unify_subject

SUBJECT_ALIASES = {"AI演習(実践)": "AI演習", "国際理解": "国際社会I:異文化理解"}   # 利用者の回答（質問13）で確定


@dataclass(frozen=True)
class FilledValue:
    student_id: str
    subject: str
    item: str
    value: Fraction
    explanation: str
    evidence: tuple[str, ...]


def _subject(name: str) -> str:
    n = unify_subject(name)
    return SUBJECT_ALIASES.get(n, n)


def _columns(ws) -> dict[tuple[str, str], int]:
    out, current = {}, None
    for c in range(5, ws.max_column + 1):
        head = ws.cell(3, c).value
        if isinstance(head, str) and head.strip():
            current = _subject(head)
        item = ws.cell(4, c).value
        if current and isinstance(item, str) and item.strip():
            out[(current, unify_subject(item))] = c
    return out


def fill_template(src: str | Path, dst: str | Path, values: Sequence[FilledValue]) -> Path:
    wb = load_workbook(src)
    ws = wb.worksheets[0]
    columns = _columns(ws)
    rows = {str(ws.cell(r, 2).value).strip(): r for r in range(6, ws.max_row + 1) if ws.cell(r, 2).value}
    if "根拠" in wb.sheetnames:
        ev = wb["根拠"]
    else:
        ev = wb.create_sheet("根拠")
        ev.append(["学籍番号", "科目", "評価項目", "セル", "値", "正確な値", "計算", "元のセル"])
        for cell in ev[1]:
            cell.font = Font(bold=True)
    written = set()
    for v in values:
        key = (_subject(v.subject), unify_subject(v.item))
        if key not in columns:
            raise KeyError(f"成績表に「{v.subject}」の「{v.item}」の列が無い")
        if v.student_id not in rows:
            raise KeyError(f"成績表に学籍番号 {v.student_id} が無い")
        r, c = rows[v.student_id], columns[key]
        if (r, c) in written:
            raise ValueError(f"{v.student_id} {v.subject} {v.item} に2回書こうとした")
        written.add((r, c))
        number = int(v.value) if v.value.denominator == 1 else float(v.value)
        ws.cell(r, c).value = number
        ev.append([v.student_id, v.subject, v.item, f"{get_column_letter(c)}{r}", number,
                   format_number(v.value).split("（")[0], v.explanation, "、".join(v.evidence)])
    for col, width in zip("ABCDEFGH", (12, 22, 12, 7, 8, 10, 40, 60)):
        ev.column_dimensions[col].width = width
    ev.freeze_panes = "A2"
    dst = Path(dst)
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    return dst


AI_FILL = PatternFill("solid", fgColor="DDEBF7")


def apply_ai_values(ai, values: Sequence[FilledValue]) -> int:
    """成績表のシート ai のうち、values のセルだけを計算値に置き換えて色を付ける。ほかのセルは触らない。"""
    columns = _columns(ai)
    rows = {str(ai.cell(r, 2).value).strip(): r for r in range(6, ai.max_row + 1) if ai.cell(r, 2).value}
    written = set()
    for v in values:
        key = (_subject(v.subject), unify_subject(v.item))
        if key not in columns:
            raise KeyError(f"成績表に「{v.subject}」の「{v.item}」の列が無い")
        if v.student_id not in rows:
            raise KeyError(f"成績表に学籍番号 {v.student_id} が無い")
        r, c = rows[v.student_id], columns[key]
        if (r, c) in written:
            raise ValueError(f"{v.student_id} {v.subject} {v.item} に2回書こうとした")
        written.add((r, c))
        ai.cell(r, c).value = int(v.value) if v.value.denominator == 1 else float(v.value)
        ai.cell(r, c).fill = AI_FILL
    return len(written)


_RATE_REF = re.compile(r"\(1-([A-Z]+)\d+\)")


def rate_columns(original) -> dict[str, int]:
    """原本の出席の式（=10-((1-I6)/0.04)）が参照している出席率の列を、科目ごとに返す。

    式が無い科目は、ほかの科目と同じ位置（出席の列＋4）を、その列の値がすべて0〜1のときだけ使う。
    """
    out = {}
    for (subject, item), col in _columns(original).items():
        if item != "出席":
            continue
        refs = {m.group(1) for r in range(6, original.max_row + 1)
                for m in [_RATE_REF.search(str(original.cell(r, col).value or ""))] if m}
        if len(refs) == 1:
            out[subject] = column_index_from_string(refs.pop())
            continue
        guess = col + 4
        values = [original.cell(r, guess).value for r in range(6, original.max_row + 1) if original.cell(r, 2).value]
        if original.cell(4, guess).value is None and values and all(isinstance(v, (int, float)) and 0 <= v <= 1 for v in values):
            out[subject] = guess
    return out


def apply_rates(ai, rate_cols: dict[str, int], rates: dict[tuple[str, str], Fraction]) -> int:
    """出席率の列を、出席簿から計算した出席率（丸めない）に置き換えて色を付ける。"""
    rows = {str(ai.cell(r, 2).value).strip(): r for r in range(6, ai.max_row + 1) if ai.cell(r, 2).value}
    n = 0
    for (sid, subject), rate in rates.items():
        col = rate_cols.get(_subject(subject))
        if col is None or sid not in rows:
            continue
        ai.cell(rows[sid], col).value = int(rate) if rate.denominator == 1 else float(rate)
        ai.cell(rows[sid], col).fill = AI_FILL
        n += 1
    return n


FIX_FILL = PatternFill("solid", fgColor="FCE4D6")   # 原本の式を利用者の指示で直したセル
_SUM = re.compile(r"^=SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)$")


def exclude_rates_from_totals(ai, rate_cols: dict[str, int]) -> list[tuple[str, str, str]]:
    """合計の式（=SUM(左:右)）が出席率の列まで足していたら、その列を外す（利用者の指示）。直したセルは色を変える。

    出席率の列が範囲の右端にあるときだけ直す。範囲の途中にあるときは、推測で直さずに止める。
    返り値: [(セル, 元の式, 直した式)]。
    """
    fixed = []
    for (subject, item), col in _columns(ai).items():
        if item != "合計" or subject not in rate_cols:
            continue
        rate = rate_cols[subject]
        for r in range(6, ai.max_row + 1):
            cell = ai.cell(r, col)
            m = _SUM.match(str(cell.value or ""))
            if not ai.cell(r, 2).value or not m:
                continue
            left, right = column_index_from_string(m.group(1)), column_index_from_string(m.group(3))
            if not left <= rate <= right:
                continue
            if rate != right:
                raise ValueError(f"{cell.coordinate} の合計 {cell.value} の途中に出席率の列がある（直し方を決められない）")
            new = f"=SUM({m.group(1)}{m.group(2)}:{get_column_letter(right - 1)}{m.group(4)})"
            fixed.append((cell.coordinate, cell.value, new))
            cell.value, cell.fill = new, FIX_FILL
    return fixed


DIFF_FILL = PatternFill("solid", fgColor="FF0000")   # 原本と値が違うセル（赤塗り・白文字）
DIFF_FONT = Font(color="FFFFFF", bold=True)


def apply_corrections(ai, corrections: dict[tuple[str, str, str], object]) -> list[str]:
    """利用者が指示した値の訂正 {(学籍番号, 科目, 評価項目): 値} を書く。セルが見つからなければ止める。"""
    columns = _columns(ai)
    rows = {str(ai.cell(r, 2).value).strip(): r for r in range(6, ai.max_row + 1) if ai.cell(r, 2).value}
    out = []
    for (sid, subject, item), value in corrections.items():
        key = (_subject(subject), unify_subject(item))
        if key not in columns or sid not in rows:
            raise KeyError(f"訂正先が見つからない: {sid} {subject} {item}")
        cell = ai.cell(rows[sid], columns[key])
        cell.value = value
        out.append(cell.coordinate)
    return out


def mark_differences(ai, ai_values, original_values) -> list[str]:
    """再計算した値で AI計算版と原本を比べ、値が違うセルを赤塗り・白文字にする（ai は式のままのシート）。"""
    out = []
    for r in range(6, max(ai_values.max_row, original_values.max_row) + 1):
        for c in range(1, max(ai_values.max_column, original_values.max_column) + 1):
            a, b = ai_values.cell(r, c).value, original_values.cell(r, c).value
            same = (a == b or (isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool)
                               and abs(a - b) <= 1e-9 * max(1.0, abs(b))))
            if not same:
                cell = ai.cell(r, c)
                cell.fill, cell.font = DIFF_FILL, DIFF_FONT
                out.append(cell.coordinate)
    return out
