"""AI計算版のシートに数式でつながった集計シートを作る（個人別評定・E一覧・テスト分析・度数分布・偏差値）。

どのセルも AI計算版のセルを参照する式なので、AI計算版を直せば全シートが自動で変わる。
使う関数は Excel・Googleスプレッドシート・LibreOffice に共通のもの（INDEX, LARGE, COUNTIF(S), AVERAGE, STDEVP, CORREL など）だけ。
"""

from collections.abc import Sequence
from dataclasses import dataclass

from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

_HEAD = PatternFill("solid", fgColor="DDE7F3")
_E = PatternFill("solid", fgColor="F8CBAD")
_VERDICT = {"適正": "E2EFDA", "易しすぎ": "FFF2CC", "難しすぎ": "F8CBAD", "差がつかない": "FCE4D6"}


@dataclass(frozen=True)
class Subject:
    name: str
    items: tuple[int, ...]        # 評価項目の列（見出しのある列だけ）
    total: int
    grade: int
    tests: tuple[tuple[int, str], ...]   # （列, 項目名）


@dataclass(frozen=True)
class Dept:
    name: str        # 国際ビジネス科
    short: str       # 国際
    sheet: str       # AI計算版_国際
    rows: tuple[int, ...]
    subjects: tuple[Subject, ...]
    gpa: int | None
    info: dict

    def ref(self, col: int, row: int) -> str:
        return f"'{self.sheet}'!${L(col)}${row}"

    def rng(self, col: int) -> str:
        return f"'{self.sheet}'!${L(col)}${self.rows[0]}:${L(col)}${self.rows[-1]}"


def describe(ws, name: str, short: str, info: dict) -> Dept:
    """AI計算版シートの並び（科目・評価項目・合計・評定・GPA・学生の行）を読む。科目の構成は値ではなく見出しで決める。"""
    rows = tuple(r for r in range(6, ws.max_row + 1) if ws.cell(r, 2).value)
    groups, current, gpa = [], None, None
    for c in range(5, ws.max_column + 1):
        head = ws.cell(3, c).value
        if isinstance(head, str) and head.strip():
            if head.strip() == "GPA":
                gpa = c
                current = None
                continue
            current = {"name": head.strip(), "items": [], "total": None, "grade": None}
            groups.append(current)
        item = ws.cell(4, c).value
        if current is None or not isinstance(item, str) or not item.strip():
            continue
        label = item.strip()
        if label == "合計":
            current["total"] = c
        elif label == "評定":
            current["grade"] = c
        elif current["total"] is None:
            current["items"].append((c, label.replace("\n", "")))
    subjects = []
    for g in groups:
        if g["total"] is None or g["grade"] is None:
            continue
        taken = any(isinstance(ws.cell(r, c).value, (int, float)) or (isinstance(ws.cell(r, c).value, str) and ws.cell(r, c).value.startswith("="))
                    for r in rows for c, _ in g["items"])
        if not taken:
            continue      # 学科の誰も受講していない科目（国際の総合ビジネス概論）
        subjects.append(Subject(g["name"], tuple(c for c, _ in g["items"]), g["total"], g["grade"],
                                tuple((c, n) for c, n in g["items"] if "テスト" in n)))
    return Dept(name, short, ws.title, rows, tuple(subjects), gpa, info)


def _bold(row):
    for cell in row:
        cell.font, cell.fill = Font(bold=True), _HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _kana(d: Dept, r: int) -> str:
    return f'=TRIM(SUBSTITUTE(SUBSTITUTE({d.ref(4, r)},CHAR(10)," "),"　"," "))'


def _taken(d: Dept, s: Subject, r: int) -> str:
    return f"COUNT({','.join(d.ref(c, r) for c in s.items)})>0"


# ---------- 個人別評定 ----------

def personal_sheet(wb, d: Dept):
    """学生ごとの各科目の評定・GPA・総点。受講していない科目は「—」。"""
    ws = wb.create_sheet(f"個人別評定_{d.short}")
    head = ["学籍番号", "氏名", "カタカナ", *[s.name for s in d.subjects], "GPA", "GPA評定", "総点"]
    ws.append(head)
    _bold(ws[1])
    for i, r in enumerate(d.rows):
        line = 2 + i
        ws.cell(line, 1, f"={d.ref(2, r)}")
        ws.cell(line, 2, f"={d.ref(3, r)}")
        ws.cell(line, 3, _kana(d, r))
        for k, s in enumerate(d.subjects):
            ws.cell(line, 4 + k, f'=IF({_taken(d, s, r)},{d.ref(s.grade, r)},"—")')
        n = 4 + len(d.subjects)
        ws.cell(line, n, f"={d.ref(d.gpa, r)}" if d.gpa else None)
        ws.cell(line, n + 2, "=" + "+".join(f"IF({_taken(d, s, r)},{d.ref(s.total, r)},0)" for s in d.subjects))
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 12, 30, 24
    ws.freeze_panes = "D2"
    return ws


def personal_layout(d: Dept) -> dict:
    """個人別評定シートのどの列に何があるか。"""
    return {"sheet": f"個人別評定_{d.short}", "first": 2, "last": 1 + len(d.rows),
            "subject_col": {s.name: 4 + k for k, s in enumerate(d.subjects)},
            "total_col": 6 + len(d.subjects)}


# ---------- E一覧 ----------

def e_sheet(wb, depts: Sequence[Dept], title: str):
    """落ちこぼれ（評定E）の一覧。Eのある学生だけを、E科目数の多い順に式で並べる。セルはその科目の合計点。"""
    many = len(depts) > 1
    subjects = list(dict.fromkeys(s.name for d in depts for s in d.subjects))
    ws = wb.create_sheet(title)
    lead = 1 if many else 0
    first_subject = lead + 5

    # 右側の作業列（学生ごとのE科目数と並べ替えの鍵）。見えなくてよいので列を隠す
    work = first_subject + len(subjects) + 2
    students = [(d, i) for d in depts for i in range(len(d.rows))]
    n = len(students)
    ws.cell(5, work, "作業: 通し番号")
    ws.cell(5, work + 1, "E科目数")
    ws.cell(5, work + 2, "並べ替えの鍵")
    for g, (d, i) in enumerate(students, start=1):
        p = personal_layout(d)
        row = 5 + g
        rng = f"'{p['sheet']}'!$D${p['first'] + i}:${L(3 + len(d.subjects))}${p['first'] + i}"
        ws.cell(row, work, g)
        ws.cell(row, work + 1, f'=COUNTIF({rng},"E")')
        ws.cell(row, work + 2, f"=IF({L(work + 1)}{row}>0,{L(work + 1)}{row}*1000+(1000-{L(work)}{row}),0)")
    keys = f"${L(work + 2)}$6:${L(work + 2)}${5 + n}"
    for c in range(work, work + 3):
        ws.column_dimensions[L(c)].hidden = True

    def header(name, k):
        vals = {d.short: d.info.get(name, ("", ""))[k] for d in depts if any(s.name == name for s in d.subjects)}
        if len(set(vals.values())) <= 1:
            return next(iter(vals.values()), "")
        return "／".join(f"{v}（{short}）" for short, v in vals.items())

    for r, label, values in ((1, "科目", subjects), (2, "担当", [header(s, 0) for s in subjects]),
                             (3, "曜日", [header(s, 1) for s in subjects])):
        ws.cell(r, lead + 1, label)
        for k, v in enumerate(values):
            ws.cell(r, first_subject + k, v)
    ws.cell(4, lead + 1, "Eの人数")
    ws.cell(4, lead + 4, "=" + "+".join(f'COUNTIF({_grade_range(d, s.name)},"E")' for d in depts for s in d.subjects))
    for k, name in enumerate(subjects):
        parts = [f'COUNTIF({_grade_range(d, name)},"E")' for d in depts if any(s.name == name for s in d.subjects)]
        ws.cell(4, first_subject + k, "=" + "+".join(parts))
    for c, v in enumerate([*(["学科"] if many else []), "学籍番号", "氏名", "カタカナ", "E科目数", *["合計点"] * len(subjects)], 1):
        ws.cell(5, c, v)
    for r in range(1, 6):
        _bold(ws[r][: first_subject + len(subjects) - 1])

    for k in range(1, n + 1):
        row = 5 + k
        g = f"(1000-MOD(LARGE({keys},{k}),1000))"
        alive = f"LARGE({keys},{k})>0"
        offset = 0
        cells = []
        for d in depts:
            p = personal_layout(d)
            cells.append((d, p, offset))
            offset += len(d.rows)

        def pick(value_for):
            expr = '""'
            for d, p, off in reversed(cells):
                expr = f"IF({g}<={off + len(d.rows)},{value_for(d, p, off)},{expr})"
            return f"=IF({alive},{expr},\"\")"

        if many:
            ws.cell(row, 1, pick(lambda d, p, off: f'"{d.name}"'))
        ws.cell(row, lead + 1, pick(lambda d, p, off: f"INDEX('{p['sheet']}'!$A:$A,{g}-{off}+1)"))
        ws.cell(row, lead + 2, pick(lambda d, p, off: f"INDEX('{p['sheet']}'!$B:$B,{g}-{off}+1)"))
        ws.cell(row, lead + 3, pick(lambda d, p, off: f"INDEX('{p['sheet']}'!$C:$C,{g}-{off}+1)"))
        ws.cell(row, lead + 4, f'=IF({alive},INT(LARGE({keys},{k})/1000),"")')
        for j, name in enumerate(subjects):
            def val(d, p, off, name=name):
                s = next((s for s in d.subjects if s.name == name), None)
                if s is None:
                    return '""'
                idx = f"{g}-{off}"
                grade = f"INDEX('{p['sheet']}'!${L(p['subject_col'][name])}:${L(p['subject_col'][name])},{idx}+1)"
                total = f"INDEX({d.rng(s.total)},{idx})"
                return f'IF({grade}="E",{total},"")'
            ws.cell(row, first_subject + j, pick(val))
    last = 5 + n
    area = f"{L(first_subject)}6:{L(first_subject + len(subjects) - 1)}{last}"
    ws.conditional_formatting.add(area, FormulaRule(formula=[f'LEN({L(first_subject)}6)>0'], fill=_E))
    for i, w in enumerate([*([14] if many else []), 12, 30, 24, 8], 1):
        ws.column_dimensions[L(i)].width = w
    ws.freeze_panes = f"{L(first_subject)}6"
    return ws


def _grade_range(d: Dept, name: str) -> str:
    p = personal_layout(d)
    col = L(p["subject_col"][name])
    return f"'{p['sheet']}'!${col}${p['first']}:${col}${p['last']}"


# ---------- テスト分析・偏差値 ----------

def _tests(depts):
    return [(d, s, c, n) for d in depts for s in d.subjects for c, n in s.tests]


def helper_sheet(wb, depts: Sequence[Dept], title: str = "分析用"):
    """他テストとの相関に使う作業シート。学生ごとに、各テストを除いた他科目テストの平均得点率を式で持つ。"""
    ws = wb.create_sheet(title)
    ws.sheet_state = "hidden"
    tests = _tests(depts)
    col = 1
    where = {}
    for d in depts:
        ws.cell(1, col, d.name)
        dt = [(s, c, n) for dd, s, c, n in tests if dd is d]
        for j, (s, c, n) in enumerate(dt):
            ws.cell(2, col + j, f"{s.name}:{n}")
            for i, r in enumerate(d.rows):
                others = [(c2, s2) for s2, c2, _ in dt if s2.name != s.name]
                rates = ",".join(f'IF(ISNUMBER({d.ref(c2, r)}),{d.ref(c2, r)}/{d.ref(c2, 5)},"")' for c2, _ in others)
                ws.cell(3 + i, col + j, f'=IFERROR(AVERAGE({rates}),"")')
            where[(d.short, s.name, n)] = f"'{title}'!${L(col + j)}$3:${L(col + j)}${2 + len(d.rows)}"
        col += len(dt) + 1
    return where


def analysis_sheet(wb, depts: Sequence[Dept], others: dict, title: str = "テスト分析"):
    ws = wb.create_sheet(title)
    ws.append(["テストの適正度（設問別の得点が無いため、テスト全体の得点で判定）。すべて AI計算版への式"])
    ws.append(["判定の目安: 易しすぎ＝平均得点率85%以上か天井効果（平均＋標準偏差＞満点）／難しすぎ＝平均得点率40%未満か床効果（平均−標準偏差＜0）"
               "／差がつかない＝標準偏差が満点の10%未満。他テストとの相関が低い（0.3未満）テストは、ほかの力と違うものを測っている可能性がある"])
    head = ["学科", "科目", "テスト", "判定", "受験者数", "満点", "平均", "平均得点率", "標準偏差", "最高", "最低", "中央値",
            "満点者数", "0点者数", "他テストとの相関", *[f"{i * 10}〜{i * 10 + 10}%" for i in range(10)], "満点超え（異常値）"]
    ws.append(head)
    _bold(ws[3])
    rows = []
    names = list(dict.fromkeys((s.name, n) for d in depts for s in d.subjects for c, n in s.tests))
    for name, item in names:
        found = [(d, s, c) for d in depts for s in d.subjects for c, n in s.tests if (s.name, n) == (name, item)]
        if len(found) > 1:
            rows.append(("両学科", name, item, found, None))
        for d, s, c in found:
            rows.append((d.name, name, item, [(d, s, c)], others.get((d.short, name, item))))
    for dept, name, item, found, other in rows:
        r = ws.max_row + 1
        rng = ",".join(d.rng(c) for d, s, c in found)
        maxi = found[0][0].ref(found[0][2], 5)

        def cnt(cond):
            return "+".join(f'COUNTIFS({d.rng(c)},{cond})' if "," in cond else f'COUNTIF({d.rng(c)},{cond})' for d, s, c in found)

        ws.cell(r, 1, dept)
        ws.cell(r, 2, name)
        ws.cell(r, 3, item)
        ws.cell(r, 5, f"=COUNT({rng})")
        ws.cell(r, 6, f"={maxi}")
        ws.cell(r, 7, f"=AVERAGE({rng})")
        ws.cell(r, 8, f"=G{r}/F{r}")
        ws.cell(r, 9, f"=STDEVP({rng})")
        ws.cell(r, 10, f"=MAX({rng})")
        ws.cell(r, 11, f"=MIN({rng})")
        ws.cell(r, 12, f"=MEDIAN({rng})")
        ws.cell(r, 13, "=" + "+".join(f'COUNTIF({d.rng(c)},">="&F{r})' for d, s, c in found))
        ws.cell(r, 14, "=" + "+".join(f"COUNTIF({d.rng(c)},0)" for d, s, c in found))
        if other:
            d, s, c = found[0]
            ws.cell(r, 15, f'=IFERROR(CORREL({d.rng(c)},{other}),"")')
        for b in range(10):
            lo = f'">="&F{r}*{b / 10}'
            hi = f'"<"&F{r}*{(b + 1) / 10}' if b < 9 else f'"<="&F{r}'
            ws.cell(r, 16 + b, "=" + "+".join(f"COUNTIFS({d.rng(c)},{lo},{d.rng(c)},{hi})" for d, s, c in found))
        ws.cell(r, 26, "=" + "+".join(f'COUNTIF({d.rng(c)},">"&F{r})' for d, s, c in found))
        ws.cell(r, 4, f'=IF(OR(H{r}>=0.85,G{r}+I{r}>F{r}),"易しすぎ",IF(OR(H{r}<0.4,G{r}-I{r}<0),"難しすぎ",'
                      f'IF(I{r}<0.1*F{r},"差がつかない","適正")))')
        for c, fmt in ((7, "0.0"), (8, "0%"), (9, "0.0"), (12, "0.0"), (15, "0.00")):
            ws.cell(r, c).number_format = fmt
    last = ws.max_row
    ws.conditional_formatting.add(f"Z4:Z{last}", FormulaRule(formula=["Z4>0"], fill=_E))
    for verdict, color in _VERDICT.items():
        ws.conditional_formatting.add(f"D4:D{last}", FormulaRule(formula=[f'D4="{verdict}"'], fill=PatternFill("solid", fgColor=color)))
    for col, width in zip("ABCD", (12, 26, 16, 12)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "D4"
    return ws


def deviation_sheet(wb, depts: Sequence[Dept], title: str = "偏差値"):
    ws = wb.create_sheet(title)
    tests = _tests(depts)
    ws.append(["学科", "学籍番号", "氏名", "カタカナ", *[f"{d.short}\n{s.name}\n{n}" for d, s, c, n in tests], "平均偏差値"])
    _bold(ws[1])
    for d in depts:
        for r in d.rows:
            line = ws.max_row + 1
            ws.cell(line, 1, d.name)
            ws.cell(line, 2, f"={d.ref(2, r)}")
            ws.cell(line, 3, f"={d.ref(3, r)}")
            ws.cell(line, 4, _kana(d, r))
            for k, (dd, s, c, n) in enumerate(tests):
                if dd is d:
                    x = d.ref(c, r)
                    ws.cell(line, 5 + k, f'=IF(ISNUMBER({x}),IFERROR(50+10*({x}-AVERAGE({d.rng(c)}))/STDEVP({d.rng(c)}),""),"")')
                    ws.cell(line, 5 + k).number_format = "0.0"
            end = L(4 + len(tests))
            ws.cell(line, 5 + len(tests), f'=IFERROR(AVERAGE(E{line}:{end}{line}),"")').number_format = "0.0"
    for col, width in zip("ABCD", (12, 12, 30, 24)):
        ws.column_dimensions[col].width = width
    ws.row_dimensions[1].height = 48
    ws.freeze_panes = "E2"
    return ws


# ---------- 度数分布 ----------

def distribution_sheet(wb, depts: Sequence[Dept], title: str = "度数分布"):
    ws = wb.create_sheet(title)
    n_sub = max(len(d.subjects) for d in depts)
    top = 100 * n_sub
    totals = [(d, personal_layout(d)) for d in depts]

    def trange(p):
        col = L(p["total_col"])
        return f"'{p['sheet']}'!${col}${p['first']}:${col}${p['last']}"

    ws.cell(1, 1, f"総点の度数分布（{n_sub}科目・{top}点満点）").font = Font(bold=True, size=12)
    for c, v in enumerate(["段階点", *[d.name for d in depts], "両学科"], 1):
        ws.cell(2, c, v)
    _bold(ws[2][: 2 + len(depts)])
    edges = list(range(0, top, 50))
    for i, e in enumerate(edges):
        r = 3 + i
        last = i == len(edges) - 1
        ws.cell(r, 1, f"{e}〜{top if last else e + 49}")
        hi = f'"<={top}"' if last else f'"<{e + 50}"'
        for k, (d, p) in enumerate(totals):
            ws.cell(r, 2 + k, f'=COUNTIFS({trange(p)},">={e}",{trange(p)},{hi})')
        ws.cell(r, 2 + len(depts), f"=SUM(B{r}:{L(1 + len(depts))}{r})")
    end = 2 + len(edges)
    _line_chart(ws, 2, end, 1 + len(depts) + 1, "総点の分布", "総点", "G2")
    r = end + 2
    ws.cell(r, 1, "総点の統計").font = Font(bold=True)
    for c, v in enumerate(["", "人数", "平均", "標準偏差", "最高", "最低"], 1):
        ws.cell(r + 1, c, v)
    _bold(ws[r + 1][:6])
    all_ranges = ",".join(trange(p) for _, p in totals)
    for k, (label, rng) in enumerate([*((d.name, trange(p)) for d, p in totals), ("両学科", all_ranges)]):
        line = r + 2 + k
        ws.cell(line, 1, label)
        for c, fn in enumerate(["COUNT", "AVERAGE", "STDEVP", "MAX", "MIN"], 2):
            ws.cell(line, c, f"={fn}({rng})")
            if fn in ("AVERAGE", "STDEVP"):
                ws.cell(line, c).number_format = "0.0"
    r = r + 4 + len(depts)
    for d in depts:
        ws.cell(r, 1, f"{d.name} 科目別の合計点の度数分布（100点満点）").font = Font(bold=True, size=12)
        ws.cell(r + 1, 1, "段階点")
        for k, s in enumerate(d.subjects):
            ws.cell(r + 1, 2 + k, s.name)
        _bold(ws[r + 1][: 1 + len(d.subjects)])
        for b in range(10):
            line = r + 2 + b
            ws.cell(line, 1, f"{b * 10}〜{100 if b == 9 else b * 10 + 9}")
            hi = '"<=100"' if b == 9 else f'"<{b * 10 + 10}"'
            for k, s in enumerate(d.subjects):
                ws.cell(line, 2 + k, f'=COUNTIFS({d.rng(s.total)},">={b * 10}",{d.rng(s.total)},{hi})')
        _line_chart(ws, r + 1, r + 11, 1 + len(d.subjects), f"{d.name} 科目別の分布", "合計点", f"{L(3 + len(d.subjects))}{r + 1}")
        r += 14
    ws.column_dimensions["A"].width = 14
    for col in "BCDE":
        ws.column_dimensions[col].width = 14
    return ws


def _line_chart(ws, head_row, last_row, last_col, title, x_title, anchor):
    chart = LineChart()
    chart.title, chart.y_axis.title, chart.x_axis.title = title, "人数", x_title
    chart.height, chart.width = 7.5, 16
    chart.add_data(Reference(ws, min_col=2, max_col=last_col, min_row=head_row, max_row=last_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=head_row + 1, max_row=last_row))
    ws.add_chart(chart, anchor)
