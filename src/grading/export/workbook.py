"""非エンジニアが検算できる成績資料（Excel）を作る。

- 「検算」シートの小計・合計・評価は Excel の計算式で書く。Excel が計算した値とシステムの値を並べ、
  一致・不一致を表示する。人は元の値と根拠のセルを見比べ、計算は Excel に任せて確かめられる。
- 値は丸めない。割り切れない値は「正確な値」の列に分数で示す。
"""

from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass, field
from fractions import Fraction
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from grading.calculation import CalcProblem, ItemValue, SubjectResult
from grading.domain.enums import DataStatus, ValueKind
from grading.export.gate import FinalizationReport
from grading.export.provenance import format_number
from grading.identity import Student
from grading.rules import Policy, RuleBook
from grading.validation import Matrix


@dataclass(frozen=True)
class EvidenceRef:
    file_name: str
    sheet_name: str | None
    cell: str | None


@dataclass(frozen=True)
class Entry:
    """1学生1科目分。result が None なら計算できなかった（problems に理由）。"""

    student_key: str
    subject: str
    values: Sequence[ItemValue]
    result: SubjectResult | None
    problems: Sequence[CalcProblem] = field(default_factory=tuple)


@dataclass(frozen=True)
class Notice:
    level: str                     # 「停止」または「要確認」
    text: str
    student_key: str | None = None
    subject: str | None = None


KIND_LABEL = {
    ValueKind.NUMBER: "数値", ValueKind.ZERO: "0点", ValueKind.BLANK: "空欄", ValueKind.NOT_SUBMITTED: "未提出",
    ValueKind.ABSENT: "欠席", ValueKind.UNGRADED: "未採点", ValueKind.NOT_APPLICABLE: "対象外", ValueKind.UNKNOWN: "不明",
}
STATUS_LABEL = {DataStatus.CONFIRMED: "確定", DataStatus.WARNING: "要確認", DataStatus.BLOCKED: "停止"}

_BOLD = Font(bold=True)
_TITLE = Font(bold=True, size=14)
_HEAD_FILL = PatternFill("solid", fgColor="DDE7F3")
_BLOCK_FILL = PatternFill("solid", fgColor="F2F2F2")
_WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
_STOP_FILL = PatternFill("solid", fgColor="F8CBAD")
_OK_FILL = PatternFill("solid", fgColor="E2EFDA")
_THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
_WRAP = Alignment(wrap_text=True, vertical="top")
_TOL = "0.000000001"


def _num(value: Fraction) -> float | int:
    return int(value) if value.denominator == 1 else float(value)


def _plain(value: Fraction) -> str:
    return format_number(value)


def _short(value: Fraction) -> str:
    """式の中で使う短い表記。割り切れない値は分数のまま。"""
    return _plain(value) if "/" not in _plain(value) else f"{value.numerator}/{value.denominator}"


def _header(ws, row: int, labels: Sequence[str]) -> None:
    for col, label in enumerate(labels, 1):
        c = ws.cell(row, col, label)
        c.font, c.fill, c.alignment = _BOLD, _HEAD_FILL, _WRAP


def _widths(ws, widths: Sequence[int]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_workbook(
    path: str | Path,
    *,
    students: Mapping[str, Student],
    rules: RuleBook,
    entries: Sequence[Entry],
    matrix: Matrix,
    report: FinalizationReport,
    locate: Callable[[int | str], EvidenceRef],
    notices: Sequence[Notice] = (),
) -> Path:
    wb = Workbook()
    guide = wb.active
    guide.title = "説明"
    check_rows = _write_check_sheet(wb.create_sheet("検算"), students, rules, entries, locate)
    _write_grades(wb.create_sheet("成績表"), students, entries, check_rows)
    _write_matrix(wb.create_sheet("完全性"), students, matrix)
    _write_evidence(wb.create_sheet("根拠"), students, entries, locate)
    _write_rules(wb.create_sheet("採点ルール"), rules)
    _write_notices(wb.create_sheet("確認事項"), students, entries, report, notices)
    _write_guide(guide, report)
    wb.move_sheet("成績表", offset=-1)
    for ws in wb.worksheets:   # 印刷しても横が切れないよう、横向き・横1ページに収める
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def problem_lines(e: "Entry") -> list[str]:
    """停止理由を、同じ理由ごとに項目をまとめて1行にする。"""
    grouped: dict[str, list[str]] = {}
    for p in e.problems:
        grouped.setdefault(p.detail, []).append(p.item or "")
    return [f"{detail}: {'、'.join(i for i in items if i)}" if any(items) else detail
            for detail, items in grouped.items()]


def _name(students: Mapping[str, Student], key: str) -> tuple[str, str, str]:
    s = students.get(key)
    return (s.student_number or "", s.name, s.class_ or "") if s else ("", f"（不明: {key}）", "")


def _write_guide(ws, report: FinalizationReport) -> None:
    _widths(ws, [110])
    if report.can_finalize:
        banner, fill = "この成績資料は【完成】です。止める理由はありません。", _OK_FILL
    else:
        banner, fill = "この成績資料は【未完成】です。下の理由が解消するまで、成績として使わないでください。", _STOP_FILL
    lines = [
        ("成績資料の見方と検算のしかた", _TITLE, None),
        (banner, _BOLD, fill),
        *[(f"・{r}", None, fill) for r in report.reasons],
        ("", None, None),
        ("■ シートの役割", _BOLD, None),
        ("成績表 … 学生×科目ごとの合計点と評価。「検算へ」を押すと、その計算の内訳に移動します。", None, None),
        ("検算 … 1人1科目ずつ、元の値・使った点数・計算式・根拠のセルを並べています。", None, None),
        ("完全性 … 学生×科目の表。OK＝確定、OK(要確認)＝確定だが確認してほしい点あり、??＝止まっている、NG＝データなし、—＝履修なし。", None, None),
        ("根拠 … すべての点数が、どの資料のどのセルから来たかの一覧。", None, None),
        ("採点ルール … 科目ごとの配点・満点・評価基準と、その根拠。", None, None),
        ("確認事項 … 人の判断が必要なこと、確認してほしいことの一覧。", None, None),
        ("", None, None),
        ("■ 検算のしかた（「検算」シート）", _BOLD, None),
        ("1. 「元の値（原文）」が、右側の「資料・シート・セル」に書かれた元資料の値と同じか、目で確かめます。", None, None),
        ("2. 「扱い」を見て、未提出や欠席などがルールどおりに扱われているか確かめます。", None, None),
        ("3. 小計・合計・評価の行は、Excel 自身が計算しています（「Excelの計算」列）。", None, None),
        ("   その隣の「システムの値」と比べた結果が「判定」列です。すべて「一致」なら、計算は正しいことになります。", None, None),
        ("4. 計算式は「計算の中身」列に、ふつうの式（例: (8+9)÷(10+10)×40）でも書いてあります。電卓でも確かめられます。", None, None),
        ("", None, None),
        ("■ 小数点について", _BOLD, None),
        ("システムは点数を丸めていません。割り切れない値は「正確な値」に分数（例: 200/3）で示しています。", None, None),
        ("評価は丸めない合計で決めています。端数の扱いで評価が変わり得る学生は「要確認」として確認事項に載せています。", None, None),
        ("最終的な小数点の調整は、人が行ってください。", None, None),
    ]
    for row, (text, font, fill) in enumerate(lines, 1):
        c = ws.cell(row, 1, text)
        c.alignment = _WRAP
        if font:
            c.font = font
        if fill:
            c.fill = fill


CHECK_HEAD = ["評価項目", "配点", "項目", "元の値（原文）", "値の種類", "扱い", "計算に使う点", "満点",
              "資料", "シート", "セル", "計算の中身", "Excelの計算", "システムの値", "判定", "正確な値"]


def _write_check_sheet(ws, students, rules: RuleBook, entries: Sequence[Entry], locate) -> dict[tuple[str, str], int]:
    _widths(ws, [14, 6, 12, 14, 9, 26, 10, 7, 22, 12, 7, 34, 11, 11, 8, 26])
    ws.freeze_panes = "A2"
    _header(ws, 1, CHECK_HEAD)
    starts: dict[tuple[str, str], int] = {}
    row = 2
    for e in entries:
        number, name, cls = _name(students, e.student_key)
        starts[(e.student_key, e.subject)] = row
        status = STATUS_LABEL[e.result.status] if e.result else "停止（計算していません）"
        title = ws.cell(row, 1, f"【{e.subject}】 {number} {name}（{cls}）  状態: {status}")
        title.font = _BOLD
        for col in range(1, len(CHECK_HEAD) + 1):
            ws.cell(row, col).fill = _BLOCK_FILL if e.result and e.result.status == DataStatus.CONFIRMED else (
                _WARN_FILL if e.result else _STOP_FILL)
        row += 1
        try:
            rule = rules.get(e.subject)
        except KeyError:
            ws.cell(row, 1, "この科目の採点ルールがありません。計算していません。")
            row += 2
            continue
        by_item = {v.item: v for v in e.values}
        component_cells, explain_terms = [], []
        for comp in rule.components:
            first = row
            pts, maxes = [], []
            for item in comp.items:
                v = by_item.get(item)
                ws.cell(row, 1, comp.name)
                ws.cell(row, 2, _num(comp.weight))
                ws.cell(row, 3, item)
                cap = rule.max_scores[item]
                if v is None:
                    ws.cell(row, 6, "値がありません（停止）").fill = _STOP_FILL
                else:
                    ws.cell(row, 4, "" if v.value.original is None else v.value.original)
                    ws.cell(row, 5, KIND_LABEL[v.value.kind])
                    ref = locate(v.evidence_id)
                    ws.cell(row, 9, ref.file_name)
                    ws.cell(row, 10, ref.sheet_name or "")
                    ws.cell(row, 11, ref.cell or "")
                    use, how = _usage(rule, v, cap)
                    ws.cell(row, 6, how).alignment = _WRAP
                    if v.status == DataStatus.BLOCKED or use == "stop":
                        ws.cell(row, 6).fill = _STOP_FILL
                    elif use is not None:
                        ws.cell(row, 7, _num(use))
                        ws.cell(row, 8, _num(cap))
                        pts.append(_plain(use))
                        maxes.append(_plain(cap))
                row += 1
            last = row - 1
            score = next((c for c in e.result.components if c.component == comp.name), None) if e.result else None
            ws.cell(row, 1, f"{comp.name} 小計").font = _BOLD
            ws.cell(row, 2, _num(comp.weight))
            if score is not None:
                ws.cell(row, 12, f"({'+'.join(pts)})÷({'+'.join(maxes)})×{_plain(comp.weight)}")
                ws.cell(row, 13, f"=SUM(G{first}:G{last})/SUM(H{first}:H{last})*B{row}")
                _system_and_judge(ws, row, score.score)
                component_cells.append(f"M{row}")
                explain_terms.append(_short(score.score))
            else:
                ws.cell(row, 12, "計算していません")
            row += 1
        if e.result:
            r = e.result
            ws.cell(row, 1, "合計").font = _BOLD
            ws.cell(row, 12, " + ".join(explain_terms))
            ws.cell(row, 13, "=" + "+".join(component_cells))
            _system_and_judge(ws, row, r.total)
            total_row = row
            row += 1
            ws.cell(row, 1, "評価").font = _BOLD
            ws.cell(row, 12, " / ".join(f"{g}: {_plain(m)}以上" for g, m in rule.grade_thresholds))
            ws.cell(row, 13, _grade_formula(f"M{total_row}", rule.grade_thresholds))
            ws.cell(row, 14, r.grade)
            ws.cell(row, 15, f'=IF(M{row}=N{row},"一致","不一致")')
            row += 1
            for note in r.notes:
                ws.cell(row, 1, "要確認").fill = _WARN_FILL
                ws.cell(row, 3, note)
                row += 1
        for line in problem_lines(e):
            ws.cell(row, 1, "停止").fill = _STOP_FILL
            ws.cell(row, 3, line)
            row += 1
        row += 1
    return starts


def _usage(rule, v: ItemValue, cap: Fraction) -> tuple[Fraction | None | str, str]:
    """(計算に使う点, 扱いの説明)。使う点が None なら除外、"stop" なら計算できない。"""
    kind = v.value.kind
    if v.status == DataStatus.BLOCKED:
        return "stop", "学生または値が未確定のため停止"
    if kind in (ValueKind.NUMBER, ValueKind.ZERO):
        number = Fraction(v.value.number)
        if not 0 <= number <= cap:
            return "stop", f"範囲外（0〜{_plain(cap)}）のため停止"
        return number, "そのまま使用"
    policy = rule.value_policies.get(kind)
    if policy == Policy.SCORE_ZERO:
        return Fraction(0), f"{KIND_LABEL[kind]} → 0点（採点ルールによる）"
    if policy == Policy.EXCLUDE:
        return None, f"{KIND_LABEL[kind]} → 計算から除外（採点ルールによる）"
    return "stop", f"{KIND_LABEL[kind]} の扱いが決まっていないため停止"


def _system_and_judge(ws, row: int, value: Fraction) -> None:
    ws.cell(row, 14, _num(value))
    ws.cell(row, 15, f'=IF(ABS(M{row}-N{row})<{_TOL},"一致","不一致")')
    ws.cell(row, 16, _plain(value))


def _grade_formula(cell: str, thresholds) -> str:
    expr = '""'
    for grade, minimum in sorted(thresholds, key=lambda t: t[1]):
        expr = f'IF({cell}>={_num(minimum)},"{grade}",{expr})'
    return "=" + expr


def _write_grades(ws, students, entries: Sequence[Entry], check_rows) -> None:
    _widths(ws, [12, 16, 8, 16, 10, 26, 7, 10, 12])
    ws.freeze_panes = "A2"
    _header(ws, 1, ["学籍番号", "氏名", "クラス", "科目", "合計", "正確な値", "評価", "状態", "内訳"])
    for row, e in enumerate(entries, 2):
        number, name, cls = _name(students, e.student_key)
        values = [number, name, cls, e.subject]
        if e.result:
            values += [_num(e.result.total), _plain(e.result.total), e.result.grade, STATUS_LABEL[e.result.status]]
        else:
            values += ["", "", "", "停止"]
        for col, v in enumerate(values, 1):
            ws.cell(row, col, v)
        state = ws.cell(row, 8)
        state.fill = _OK_FILL if state.value == "確定" else (_WARN_FILL if state.value == "要確認" else _STOP_FILL)
        link = ws.cell(row, 9, "検算へ")
        link.hyperlink = f"#検算!A{check_rows[(e.student_key, e.subject)]}"
        link.font = Font(color="0563C1", underline="single")


def _write_matrix(ws, students, matrix: Matrix) -> None:
    ws.cell(1, 1, "学生×科目（OK＝確定 / OK(要確認) / ??＝止まっている / NG＝データなし / —＝履修なし）").font = _BOLD
    _header(ws, 2, ["学籍番号", "氏名", *matrix.subjects])
    _widths(ws, [12, 16, *[12] * len(matrix.subjects)])
    for row, key in enumerate(matrix.students, 3):
        number, name, _ = _name(students, key)
        ws.cell(row, 1, number)
        ws.cell(row, 2, name)
        for col, sub in enumerate(matrix.subjects, 3):
            v = matrix.cells[key][sub]
            c = ws.cell(row, col, v.value)
            c.alignment = Alignment(horizontal="center")
            c.fill = {"OK": _OK_FILL, "OK(要確認)": _WARN_FILL, "??": _STOP_FILL, "NG": _STOP_FILL}.get(v.value, PatternFill())


def _write_evidence(ws, students, entries: Sequence[Entry], locate) -> None:
    _widths(ws, [12, 16, 14, 12, 14, 10, 24, 12, 8])
    ws.freeze_panes = "A2"
    _header(ws, 1, ["学籍番号", "氏名", "科目", "項目", "元の値（原文）", "値の種類", "資料", "シート", "セル"])
    row = 2
    for e in entries:
        number, name, _ = _name(students, e.student_key)
        for v in e.values:
            ref = locate(v.evidence_id)
            for col, x in enumerate([number, name, e.subject, v.item, v.value.original or "",
                                     KIND_LABEL[v.value.kind], ref.file_name, ref.sheet_name or "", ref.cell or ""], 1):
                ws.cell(row, col, x).border = _THIN
            row += 1


def _write_rules(ws, rules: RuleBook) -> None:
    _widths(ws, [16, 14, 8, 30, 50])
    row = 1
    for subject in rules.subjects():
        rule = rules.get(subject)
        ws.cell(row, 1, f"【{subject}】 根拠: {rule.rule_ref}").font = _BOLD
        row += 1
        _header(ws, row, ["評価項目", "配点", "", "項目（満点）", ""])
        row += 1
        for c in rule.components:
            ws.cell(row, 1, c.name)
            ws.cell(row, 2, _num(c.weight))
            ws.cell(row, 4, "、".join(f"{i}（{_plain(rule.max_scores[i])}点）" for i in c.items))
            row += 1
        ws.cell(row, 1, "評価基準")
        ws.cell(row, 4, " / ".join(f"{g}: {_plain(m)}点以上" for g, m in rule.grade_thresholds))
        row += 1
        ws.cell(row, 1, "記号の意味")
        ws.cell(row, 4, "、".join(f"「{k}」＝{KIND_LABEL[v]}" for k, v in rule.markers.items()) or "なし")
        row += 1
        ws.cell(row, 1, "数値以外の扱い")
        labels = {Policy.SCORE_ZERO: "0点", Policy.EXCLUDE: "計算から除外"}
        text = "、".join(f"{KIND_LABEL[k]}＝{labels[p]}" for k, p in rule.value_policies.items())
        ws.cell(row, 4, (text or "定めなし") + "（ここに無い扱いはすべて停止して人に確認します）")
        row += 2


def _write_notices(ws, students, entries: Sequence[Entry], report: FinalizationReport, notices: Sequence[Notice]) -> None:
    _widths(ws, [8, 12, 16, 14, 90])
    ws.freeze_panes = "A2"
    _header(ws, 1, ["区分", "学籍番号", "氏名", "科目", "内容"])
    rows: list[tuple[str, str | None, str | None, str]] = []
    rows += [("停止", None, None, f"完成にできない理由: {r}") for r in report.reasons]
    rows += [(n.level, n.student_key, n.subject, n.text) for n in notices]
    for e in entries:
        rows += [("停止", e.student_key, e.subject, line) for line in problem_lines(e)]
        if e.result:
            rows += [("要確認", e.student_key, e.subject, n) for n in e.result.notes]
    rows.sort(key=lambda r: r[0] != "停止")
    for row, (level, key, subject, text) in enumerate(rows, 2):
        number, name, _ = _name(students, key) if key else ("", "", "")
        for col, v in enumerate([level, number, name, subject or "", text], 1):
            c = ws.cell(row, col, v)
            c.alignment = _WRAP
        ws.cell(row, 1).fill = _STOP_FILL if level == "停止" else _WARN_FILL
    if not rows:
        ws.cell(2, 5, "確認事項はありません。")
