"""出席簿（月別シート）の読込と、科目ごとの出欠の数え直し。

- 月別シートだけを読む。「科目ごと」シートの集計式は使わない（8・9月を7月の科目並びで数える誤りがあるため）。
- 日付は 3行目から読む。文字（「4月 13日（月）」）、日付値、単純な式（=E3+1）に対応する。日付の無い見出しは日付なしとして扱う。
- 科目名は全角半角・康熙部首・前後の空白を揃えて突き合わせ、原文も残す。
- 記号は 欠・遅・×・空欄 だけを受け付け、それ以外は問題として報告する。
- 欠・遅の1件ごとに、どのシートのどのセルかを残す。
"""

import datetime as dt
import re
import unicodedata
from collections.abc import Mapping
from dataclasses import dataclass, field
from fractions import Fraction
from pathlib import Path

import jpholiday
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

WEEK = "月火水木金土日"
MARKS = (None, "欠", "遅", "×")
_TEXT_DATE = re.compile(r"(\d+)\s*月\s*(\d+)\s*日\s*[（(]\s*(\S)\s*[)）]")
_BLANK_DATE = re.compile(r"^\s*月\s*日\s*[（(]\s*(\S)\s*[)）]\s*$")
_WEEKDAY_ONLY = re.compile(r'^=TEXT\(|^\s*[（(]\s*\S\s*[)）]\s*$')
_PLUS = re.compile(r"^=\s*([A-Z]+)3\s*\+\s*(\d+)\s*$")


def unify_subject(name: str) -> str:
    """突き合わせ用の科目名。康熙部首（⽇→日）・全角半角・前後の空白を揃える。原文は別に残す。"""
    return unicodedata.normalize("NFKC", name).strip()


@dataclass(frozen=True)
class Session:
    sheet: str
    column: str
    date: dt.date | None
    period: object
    subject: str
    subject_original: str


@dataclass(frozen=True)
class Mark:
    student_id: str
    session: int          # sessions の番号
    value: str | None
    cell: str             # 「シート名!F6」


@dataclass(frozen=True)
class Problem:
    cell: str
    detail: str


@dataclass
class Register:
    path: str
    students: dict[str, tuple[str, str]] = field(default_factory=dict)   # 学籍番号 -> (氏名, 日本名)
    sessions: list[Session] = field(default_factory=list)
    marks: list[Mark] = field(default_factory=list)
    problems: list[Problem] = field(default_factory=list)
    dates: set[dt.date] = field(default_factory=set)        # 見出しに書かれていた日付


def _month_sheets(wb, prefix: str = ""):
    return [ws for ws in wb.worksheets if re.search(r"\d+月$", ws.title.strip()) and "原紙" not in ws.title
            and ws.title.startswith(prefix)]


def _end_column(ws) -> int:
    for c in ws[3]:
        if c.value == "授業合計数":
            return c.column
    raise ValueError(f"{ws.title}: 「授業合計数」の見出しが無い")


def read_register(path: str | Path, date_corrections: Mapping[str, dt.date] | None = None,
                  sheet_prefix: str = "") -> Register:
    """date_corrections: {「シート名!H3」: 正しい日付}。利用者の回答で決まった訂正だけを渡す。
    sheet_prefix: 2学科の出席簿を写したブック（成績表のAI計算版）から1学科分だけ読むときのシート名の頭。"""
    corrections = dict(date_corrections or {})
    wb = load_workbook(path)
    reg = Register(str(path))
    for ws in _month_sheets(wb, sheet_prefix):
        month = int(re.search(r"(\d+)月$", ws.title.strip()).group(1))
        end = _end_column(ws)
        rows = [r for r in range(6, ws.max_row + 1) if ws.cell(r, 2).value]
        for r in rows:
            sid = str(ws.cell(r, 2).value).strip()
            names = (ws.cell(r, 3).value, ws.cell(r, 4).value)
            if sid in reg.students and reg.students[sid] != names:
                reg.problems.append(Problem(f"{ws.title}!B{r}", f"{sid} の氏名が他のシートと違う: {reg.students[sid]} / {names}"))
            reg.students.setdefault(sid, names)

        resolved: dict[str, dt.date] = {}
        current: dt.date | None = None
        for c in range(5, end):
            letter = get_column_letter(c)
            ref = f"{ws.title}!{letter}3"
            head = ws.cell(3, c).value
            if head is not None and not (isinstance(head, str) and _WEEKDAY_ONLY.search(head)):
                current = _read_date(ws, head, ref, month, resolved, reg.problems)
                if ref in corrections:
                    current = corrections[ref]
                    reg.problems[:] = [p for p in reg.problems if p.cell != ref]
                if current is not None:
                    resolved[letter] = current
                    reg.dates.add(current)
            subject = ws.cell(4, c).value
            if subject is None:
                if any(ws.cell(r, c).value in ("欠", "遅") for r in rows):
                    reg.problems.append(Problem(f"{ws.title}!{letter}4", "科目名が空欄なのに欠・遅の記入がある"))
                continue
            index = len(reg.sessions)
            reg.sessions.append(Session(ws.title, letter, current, ws.cell(5, c).value,
                                        unify_subject(subject), subject))
            if current is None and any(ws.cell(r, c).value in ("欠", "遅") for r in rows):
                reg.problems.append(Problem(f"{ws.title}!{letter}3", "日付が無い列に欠・遅の記入がある"))
            for r in rows:
                value = ws.cell(r, c).value
                cell = f"{ws.title}!{letter}{r}"
                if value not in MARKS:
                    reg.problems.append(Problem(cell, f"想定外の記号 {value!r}"))
                reg.marks.append(Mark(str(ws.cell(r, 2).value).strip(), index, value, cell))
    return reg


def _read_date(ws, head, ref, month, resolved, problems) -> dt.date | None:
    if isinstance(head, dt.datetime):
        return head.date()
    if isinstance(head, dt.date):
        return head
    text = str(head)
    plus = _PLUS.match(text)
    if plus:
        base = resolved.get(plus.group(1))
        if base is None:
            problems.append(Problem(ref, f"日付の式 {text} の元の日付が読めない"))
            return None
        return base + dt.timedelta(days=int(plus.group(2)))
    folded = unicodedata.normalize("NFKC", text)
    if _BLANK_DATE.match(folded):
        problems.append(Problem(ref, f"日付が入っていない見出し {text.strip()!r}（授業日として数えない）"))
        return None
    m = _TEXT_DATE.search(folded)
    if not m:
        problems.append(Problem(ref, f"日付が読めない {text!r}"))
        return None
    date = dt.date(2026, int(m.group(1)), int(m.group(2)))
    written = m.group(3)
    if WEEK[date.weekday()] != written:
        problems.append(Problem(ref, f"{date:%m/%d} は{WEEK[date.weekday()]}曜日だが「{written}」と書かれている"
                                     f"（{'日曜' if date.weekday() == 6 else '土曜' if date.weekday() == 5 else '曜日違い'}）"))
    if date.month != month:
        problems.append(Problem(ref, f"{month}月のシートに {date:%m/%d} と書かれている"))
    return date


@dataclass(frozen=True)
class SubjectCount:
    student_id: str
    subject: str
    sessions: int
    absent: int
    late: int
    absent_cells: tuple[str, ...]
    late_cells: tuple[str, ...]

    @property
    def rate(self) -> Fraction:
        """出席簿の式と同じ定義：1 −（欠席 ＋ 遅刻÷3）÷ 授業数。丸めない。"""
        return 1 - (self.absent + Fraction(self.late, 3)) / self.sessions


def subject_counts(reg: Register) -> dict[tuple[str, str], SubjectCount]:
    acc: dict[tuple[str, str], list] = {}
    for m in reg.marks:
        if m.value == "×":
            continue
        s = reg.sessions[m.session]
        if s.date is None:
            continue
        a = acc.setdefault((m.student_id, s.subject), [0, [], []])
        a[0] += 1
        if m.value == "欠":
            a[1].append(m.cell)
        elif m.value == "遅":
            a[2].append(m.cell)
    return {k: SubjectCount(k[0], k[1], n, len(ab), len(lt), tuple(ab), tuple(lt)) for k, (n, ab, lt) in acc.items()}


@dataclass(frozen=True)
class CalendarAnomaly:
    date: dt.date
    kind: str
    detail: str


def calendar_anomalies(reg: Register, first: dt.date, last: dt.date) -> list[CalendarAnomaly]:
    """土日祝日は休み。それ以外の日で授業が無い・全員×、または休みの日に授業がある日を挙げる。"""
    by_date: dict[dt.date, list[int]] = {}
    for i, s in enumerate(reg.sessions):
        if s.date is not None:
            by_date.setdefault(s.date, []).append(i)
    live: dict[int, int] = {}
    for m in reg.marks:
        if m.value != "×":
            live[m.session] = live.get(m.session, 0) + 1
    out = []
    d = first
    while d <= last:
        holiday = jpholiday.is_holiday_name(d)
        weekend = d.weekday() >= 5
        idx = by_date.get(d, [])
        held = [i for i in idx if live.get(i)]
        label = f"{d:%m/%d}（{WEEK[d.weekday()]}）"
        if (holiday or weekend) and held:
            out.append(CalendarAnomaly(d, "祝日に授業の記録" if holiday else "土日に授業の記録",
                                       f"{label}{' ' + holiday if holiday else ''}: {len(held)}コマに記入"))
        elif not (holiday or weekend):
            if not idx and d in reg.dates:
                out.append(CalendarAnomaly(d, "平日なのに科目の記入が無い", label))
            elif not idx:
                out.append(CalendarAnomaly(d, "平日なのに出席簿に日付が無い", label))
            elif not held:
                subjects = sorted({reg.sessions[i].subject for i in idx})
                out.append(CalendarAnomaly(d, "平日なのに全員×", f"{label}: {len(idx)}コマ（{'、'.join(subjects)}）"))
        d += dt.timedelta(days=1)
    return out


def weekly_counts(reg: Register, subject: str, exclude: set[str] = frozenset()) -> dict[tuple[str, str], SubjectCount]:
    """週ごとに1回と数える科目（日本語運用力強化演習）の集計。

    その週の×でない授業のうち、1つでも出席なら「出席」、出席が無く遅刻があれば「遅刻」、どちらも無ければ「欠席」。
    ×しか無い週は数えない。exclude には数えない授業を「シート名!列」で渡す。
    根拠のセルは、欠席の週はその週の欠のセルすべて、遅刻の週はその週の遅のセルすべて。
    """
    weeks: dict[tuple[str, dt.date], list[Mark]] = {}
    for m in reg.marks:
        s = reg.sessions[m.session]
        if s.subject != subject or s.date is None or m.value == "×" or f"{s.sheet}!{s.column}" in exclude:
            continue
        monday = s.date - dt.timedelta(days=s.date.weekday())
        weeks.setdefault((m.student_id, monday), []).append(m)
    acc: dict[str, list] = {}
    for (sid, _), marks in sorted(weeks.items()):
        a = acc.setdefault(sid, [0, [], []])
        a[0] += 1
        values = [m.value for m in marks]
        if None in values:
            continue
        if "遅" in values:
            a[2].extend(m.cell for m in marks if m.value == "遅")
            continue
        a[1].extend(m.cell for m in marks if m.value == "欠")
    out = {}
    for sid, (n, ab, lt) in acc.items():
        absent_weeks = sum(1 for (s, _), ms in weeks.items() if s == sid and all(m.value == "欠" for m in ms))
        late_weeks = sum(1 for (s, _), ms in weeks.items()
                         if s == sid and None not in [m.value for m in ms] and "遅" in [m.value for m in ms])
        out[(sid, subject)] = SubjectCount(sid, subject, n, absent_weeks, late_weeks, tuple(ab), tuple(lt))
    return out
