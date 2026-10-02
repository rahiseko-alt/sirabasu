"""出来上がった成績ファイルの検算。生成の最後に必ず実行し、1件でも食い違えば不合格にする。

生成の処理とは別の道筋で数え直して照合する:
1. 再計算エラー（#DIV/0! など）が1つも無い
2. 原本シートが元の成績表と、式・値とも一致する
3. AI計算版のうち、AIが書いてよいセル（出席点・態度点・出席率）以外は原本と同じ
4. 出席点・態度点・出席率を、出席簿のセルを直接数え直して整数の計算で求め、AI計算版と一致する
5. 個人別評定・E一覧・テスト分析・偏差値の式の結果が、AI計算版の値から Python で計算した結果と一致する
"""

import datetime as dt
import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from fractions import Fraction
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from grading.analysis.sheets import collect_tests
from grading.analysis.test_stats import deviation_scores, judge, summarize
from grading.export.e_list import grade_rows, gpa_of
from grading.export.fill import _columns, _subject, rate_columns
from grading.importing.attendance import read_register, unify_subject

TOLERANCE = 1e-9
_SUM = re.compile(r"^=SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)$")


@dataclass(frozen=True)
class DeptSpec:
    name: str              # 国際ビジネス科
    short: str             # 国際
    register: Path         # 出席簿
    original: Path         # 元の成績表


@dataclass(frozen=True)
class Rules:
    attendance: Mapping[str, tuple[int, int]]      # 科目: (満点, 4%ごとの減点)
    weekly: str                                    # 1週＝1回と数える科目
    late_deduct: Sequence[str]                     # 態度点＝10−2×遅刻
    flat_ten: Sequence[str]                        # 態度点＝10（名前つき記録0件）
    date_corrections: Mapping[str, dt.date] = field(default_factory=dict)
    corrections: Mapping[str, Mapping[tuple[str, str, str], object]] = field(default_factory=dict)  # 学科: {(学籍番号, 科目, 項目): 値}


@dataclass(frozen=True)
class Finding:
    check: str
    where: str
    detail: str


@dataclass
class Report:
    compared: dict[str, int] = field(default_factory=dict)
    findings: list[Finding] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.findings

    def add(self, check: str, n: int = 1):
        self.compared[check] = self.compared.get(check, 0) + n

    def text(self) -> str:
        lines = ["検算結果: " + ("合格（食い違い0件）" if self.ok else f"不合格（食い違い{len(self.findings)}件）"), ""]
        for check, n in self.compared.items():
            bad = sum(1 for f in self.findings if f.check == check)
            lines.append(f"{'○' if not bad else '×'} {check}: {n}件を照合、食い違い{bad}件")
        if self.findings:
            lines += ["", "食い違いの一覧:"]
            lines += [f"- [{f.check}] {f.where}: {f.detail}" for f in self.findings]
        return "\n".join(lines) + "\n"


def _same(a, b) -> bool:
    if isinstance(a, bool) or isinstance(b, bool):
        return a == b
    if isinstance(a, (int, float, Fraction)) and isinstance(b, (int, float, Fraction)):
        return abs(float(a) - float(b)) <= TOLERANCE * max(1.0, abs(float(b)))
    return (a in (None, "") and b in (None, "")) or a == b


# ---------- 出席簿を直接数え直す（生成側の集計処理は使わない） ----------

def raw_counts(register: Path) -> dict[tuple[str, str], tuple[int, int, int]]:
    """月別シートのセルを直接数える。{(学籍番号, 科目): (授業数, 欠, 遅)}。× は授業数に入れない。"""
    wb = load_workbook(register)
    out: dict[tuple[str, str], list[int]] = {}
    for ws in wb.worksheets:
        if not re.search(r"\d+月$", ws.title.strip()) or "原紙" in ws.title:
            continue
        end = next(c.column for c in ws[3] if c.value == "授業合計数")
        rows = [r for r in range(6, ws.max_row + 1) if ws.cell(r, 2).value]
        for col in range(5, end):
            subject = ws.cell(4, col).value
            if subject is None:
                continue
            for r in rows:
                v = ws.cell(r, col).value
                a = out.setdefault((str(ws.cell(r, 2).value).strip(), unify_subject(subject)), [0, 0, 0])
                a[0] += v != "×"
                a[1] += v == "欠"
                a[2] += v == "遅"
    return {k: tuple(v) for k, v in out.items()}


def raw_weekly(register: Path, subject: str, corrections) -> dict[str, tuple[int, int, int]]:
    """1週＝1回の科目。その週に1回でも出席なら出席、出席が無く遅があれば遅刻、ほかは欠席。×だけの週は数えない。"""
    reg = read_register(register, date_corrections=corrections)
    weeks: dict[tuple[str, dt.date], list] = {}
    for m in reg.marks:
        s = reg.sessions[m.session]
        if s.subject == subject and s.date is not None and m.value != "×":
            weeks.setdefault((m.student_id, s.date - dt.timedelta(days=s.date.weekday())), []).append(m.value)
    out: dict[str, list[int]] = {}
    for (sid, _), values in weeks.items():
        a = out.setdefault(sid, [0, 0, 0])
        a[0] += 1
        if None not in values:
            a[2 if "遅" in values else 1] += 1
    return {k: tuple(v) for k, v in out.items()}


def expected_points(sessions: int, absent: int, late: int, maximum: int, step: int) -> int:
    """整数だけで計算する出席点。出席率＝(3×授業数−3×欠−遅)÷(3×授業数)。60%未満は0点、4%下がるごとに減点。"""
    whole, lost = 3 * sessions, 3 * absent + late
    if 5 * (whole - lost) < 3 * whole:
        return 0
    return maximum - step * ((100 * lost) // (4 * whole))


# ---------- 照合 ----------

def _rows(ws) -> dict[str, int]:
    return {str(ws.cell(r, 2).value).strip(): r for r in range(6, ws.max_row + 1) if ws.cell(r, 2).value}


def check_errors(values_wb, report: Report):
    for ws in values_wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                report.add("再計算エラーが無いか")
                if isinstance(c.value, str) and (c.value.startswith("#") or c.value.startswith("Err:")):
                    report.findings.append(Finding("再計算エラーが無いか", f"{ws.title}!{c.coordinate}", c.value))


def check_original(out_ws, original_ws, report: Report, check: str = "原本シートが元の成績表と同じか"):
    for r in range(1, max(out_ws.max_row, original_ws.max_row) + 1):
        for c in range(1, max(out_ws.max_column, original_ws.max_column) + 1):
            report.add(check)
            a, b = out_ws.cell(r, c).value, original_ws.cell(r, c).value
            if not _same(a, b):
                report.findings.append(Finding(check, f"{out_ws.title}!{out_ws.cell(r, c).coordinate}", f"{a!r} ≠ 元 {b!r}"))


def check_red(ai_ws, ai_values_ws, original_values_ws, report: Report):
    """原本と値が違うセルは赤、同じセルは赤でないこと。"""
    check = "原本と値が違うセルだけが赤か"
    for r in range(6, max(ai_values_ws.max_row, original_values_ws.max_row) + 1):
        for c in range(1, max(ai_values_ws.max_column, original_values_ws.max_column) + 1):
            differs = not _same(ai_values_ws.cell(r, c).value, original_values_ws.cell(r, c).value)
            red = str(ai_ws.cell(r, c).fill.fgColor.rgb).endswith("FF0000")
            report.add(check)
            if differs != red:
                report.findings.append(Finding(check, f"{ai_ws.title}!{ai_ws.cell(r, c).coordinate}",
                                               "原本と違うのに赤でない" if differs else "原本と同じなのに赤"))


def ai_cells(original_ws, rules: Rules) -> set[tuple[int, int]]:
    """AIが書いてよいセル。"""
    columns = _columns(original_ws)
    cols = {columns[(_subject(s), "出席")] for s in rules.attendance}
    cols |= {columns[(_subject(s), "授業態度")] for s in [*rules.late_deduct, *rules.flat_ten]}
    rates = rate_columns(original_ws)
    cols |= {rates[_subject(s)] for s in rules.attendance if _subject(s) in rates}
    return {(r, c) for r in _rows(original_ws).values() for c in cols}


def check_untouched(ai_ws, original_ws, rules: Rules, report: Report, corrections=None):
    check = "AIが書かないセルは原本のままか"
    allowed = ai_cells(original_ws, rules)
    columns, rows = _columns(original_ws), _rows(original_ws)
    corrected = set()
    for (sid, subject, item), value in (corrections or {}).items():
        r, c = rows[sid], columns[(_subject(subject), unify_subject(item))]
        corrected.add((r, c))
        report.add("指示された訂正が入っているか")
        if not _same(ai_ws.cell(r, c).value, value):
            report.findings.append(Finding("指示された訂正が入っているか", f"{ai_ws.title}!{ai_ws.cell(r, c).coordinate}",
                                           f"{ai_ws.cell(r, c).value!r}、指示は {value!r}"))
    rate_cols = set(rate_columns(original_ws).values())
    fixed = {c for s in rules.flat_ten for c in [_columns(original_ws)[(_subject(s), "授業態度")]]}
    for r, c in sorted(allowed):
        if c in fixed:
            continue
        report.add("AIのセルが数字の貼り付けでなく式か")
        v = ai_ws.cell(r, c).value
        if not (isinstance(v, str) and v.startswith("=")):
            report.findings.append(Finding("AIのセルが数字の貼り付けでなく式か",
                                           f"{ai_ws.title}!{ai_ws.cell(r, c).coordinate}", f"式ではなく {v!r}"))
    for r in range(1, max(ai_ws.max_row, original_ws.max_row) + 1):
        for c in range(1, max(ai_ws.max_column, original_ws.max_column) + 1):
            if (r, c) in allowed or (r, c) in corrected:
                continue
            report.add(check)
            a, b = ai_ws.cell(r, c).value, original_ws.cell(r, c).value
            if (r, c) == (1, 1) and isinstance(b, str):
                b += "（AI計算版）"
            total = _SUM.match(str(b or ""))
            if total and r >= 6 and column_index_from_string(total.group(3)) in rate_cols:
                # 合計が出席率の列まで足していた式は、その列を外す（利用者の指示）
                b = f"=SUM({total.group(1)}{total.group(2)}:{get_column_letter(column_index_from_string(total.group(3)) - 1)}{total.group(4)})"
                report.add("合計から出席率の列を外したか")
                if not _same(a, b):
                    report.findings.append(Finding("合計から出席率の列を外したか",
                                                   f"{ai_ws.title}!{ai_ws.cell(r, c).coordinate}", f"{a!r}、正しくは {b!r}"))
                continue
            if not _same(a, b):
                report.findings.append(Finding(check, f"{ai_ws.title}!{ai_ws.cell(r, c).coordinate}", f"{a!r} ≠ 原本 {b!r}"))


def check_ai_values(ai_values_ws, original_values_ws, spec: DeptSpec, rules: Rules, report: Report):
    """出席点・態度点・出席率を出席簿から数え直して照合する。樋口先生の態度点は原本と同じであること。"""
    counts = raw_counts(spec.register)
    weekly = raw_weekly(spec.register, rules.weekly, rules.date_corrections)
    columns, rates, rows = _columns(ai_values_ws), rate_columns(original_values_ws), _rows(ai_values_ws)
    for sid, r in rows.items():
        for subject, (maximum, step) in rules.attendance.items():
            n, absent, late = weekly[sid] if subject == rules.weekly else counts[(sid, subject)]
            where = f"{ai_values_ws.title} {sid} {subject}"
            got = ai_values_ws.cell(r, columns[(_subject(subject), "出席")]).value
            want = expected_points(n, absent, late, maximum, step)
            report.add("出席点（出席簿から数え直し）")
            if not _same(got, want):
                report.findings.append(Finding("出席点（出席簿から数え直し）", where,
                                               f"表は{got!r}、数え直すと{want}（授業{n}・欠{absent}・遅{late}）"))
            if _subject(subject) in rates:
                report.add("出席率（出席簿から数え直し）")
                got = ai_values_ws.cell(r, rates[_subject(subject)]).value
                want = Fraction(3 * n - 3 * absent - late, 3 * n)
                if not _same(got, want):
                    report.findings.append(Finding("出席率（出席簿から数え直し）", where, f"表は{got!r}、数え直すと{float(want)}"))
            attitude = columns[(_subject(subject), "授業態度")]
            got = ai_values_ws.cell(r, attitude).value
            if subject in rules.late_deduct:
                want, how = 10 - 2 * counts[(sid, subject)][2], "10−2×遅刻"
            elif subject in rules.flat_ten:
                want, how = 10, "10"
            else:
                want, how = original_values_ws.cell(r, attitude).value, "原本のまま"
            check = "態度点（AI計算分は出席簿から数え直し、ほかは原本のまま）"
            report.add(check)
            if not _same(got, want):
                report.findings.append(Finding(check, where, f"表は{got!r}、{how}なら{want!r}"))


def check_linked(values_wb, depts: Sequence[DeptSpec], report: Report):
    """式でつないだシートの結果を、AI計算版の値から Python で計算した結果と照合する。"""
    rows_of, tests = {}, []
    for d in depts:
        ai = values_wb[f"AI計算版_{d.short}"]
        rows = grade_rows(ai)
        rows_of[d.name] = rows
        tests += collect_tests(ai, d.name)
        _check_personal(values_wb[f"個人別評定_{d.short}"], rows, gpa_of(ai), report)
        _check_e(values_wb[f"E一覧_{d.short}"], [d], rows_of, report)
    _check_e(values_wb["E一覧_統合"], depts, rows_of, report)
    _check_analysis(values_wb["テスト分析"], tests, report)
    _check_deviation(values_wb["偏差値"], depts, tests, report)


def _check_personal(ws, rows, gpa, report: Report):
    check = "個人別評定（評定・GPA・総点）"
    head = [c.value for c in ws[1]]
    grade = {(r.student_id, r.subject): r.grade for r in rows}
    total: dict[str, float] = {}
    for r in rows:
        if isinstance(r.total, (int, float)):
            total[r.student_id] = total.get(r.student_id, 0) + r.total
    for line in ws.iter_rows(min_row=2, values_only=True):
        sid = line[0]
        if not sid:
            continue
        pairs = [(h, line[3 + k], grade.get((sid, h), "—")) for k, h in enumerate(head[3:-3])]
        pairs += [("GPA", line[-3], gpa.get(sid)), ("総点", line[-1], total.get(sid, 0))]
        for label, got, want in pairs:
            report.add(check)
            if not _same(got, want):
                report.findings.append(Finding(check, f"{ws.title} {sid} {label}", f"表は{got!r}、計算では{want!r}"))


def _check_e(ws, depts, rows_of, report: Report):
    check = "E一覧（Eの学生・科目・点数・並び順）"
    lead = 1 if len(depts) > 1 else 0
    subjects = [c.value for c in ws[1]][lead + 4:]
    subjects = subjects[:subjects.index(None)] if None in subjects else subjects
    n_students = sum(len({r.student_id for r in rows_of[d.name]}) for d in depts)
    got, counts = {}, []
    for line in ws.iter_rows(min_row=6, max_row=5 + n_students, values_only=True):
        sid = line[lead]
        if not sid:
            continue
        counts.append(line[lead + 3])
        for k, s in enumerate(subjects):
            if line[lead + 4 + k] not in (None, ""):
                got[(sid, s)] = line[lead + 4 + k]
    want = {(r.student_id, r.subject): r.total for d in depts for r in rows_of[d.name] if r.grade == "E"}
    report.add(check, len(want) + 1)
    for key in sorted(set(got) | set(want)):
        if key not in got or key not in want or not _same(got[key], want[key]):
            report.findings.append(Finding(check, f"{ws.title} {key[0]} {key[1]}",
                                           f"表は{got.get(key, '（無し）')!r}、計算では{want.get(key, '（Eではない）')!r}"))
    if counts != sorted(counts, reverse=True):
        report.findings.append(Finding(check, ws.title, "E科目数の多い順に並んでいない"))


_ANALYSIS = ["判定", "受験者数", "満点", "平均", "平均得点率", "標準偏差", "最高", "最低", "中央値", "満点の人数", "0点の人数"]


def _check_analysis(ws, tests, report: Report):
    check = "テスト分析（判定・平均・標準偏差・分布）"
    for line in ws.iter_rows(min_row=4, values_only=True):
        dept, subject, item = line[0], line[1], line[2]
        if not subject:
            continue
        ts = [t for t in tests if t.subject == subject and t.item == item and (dept == "両学科" or t.dept == dept)]
        where = f"{ws.title} {dept} {subject} {item}"
        if not ts:
            report.findings.append(Finding(check, where, "AI計算版にこのテストが見つからない"))
            continue
        maximum = ts[0].maximum
        scores = [v for t in ts for v in t.scores.values()]
        s = summarize(scores, maximum)
        want = [judge(s)[0], s.n, s.maximum, s.mean, s.rate, s.sd, s.maximum_score, s.minimum, s.median, s.full_marks, s.zeros]
        inside = summarize([v for v in scores if v <= maximum], maximum)
        pairs = list(zip(_ANALYSIS, line[3:14], want))
        pairs += [("得点分布", list(line[15:25]), list(inside.bins)), ("満点超え", line[25], len(scores) - inside.n)]
        for label, got, w in pairs:
            report.add(check)
            if isinstance(w, list):
                bad = len(got) != len(w) or not all(_same(a, b) for a, b in zip(got, w))
            else:
                bad = not _same(got, w)
            if bad:
                report.findings.append(Finding(check, f"{where} {label}", f"表は{got!r}、計算では{w!r}"))


def _check_deviation(ws, depts, tests, report: Report):
    check = "偏差値"
    short_to_name = {d.short: d.name for d in depts}
    head = [c.value for c in ws[1]]
    columns = []
    for k, h in enumerate(head[4:-1]):
        short, subject, item = h.split("\n", 2)
        item = item.replace("\n", "")
        found = [t for t in tests if t.dept == short_to_name[short] and t.subject == subject and t.item == item]
        columns.append((4 + k, short_to_name[short], found[0] if found else None, h))
    for line in ws.iter_rows(min_row=2, values_only=True):
        dept, sid = line[0], line[1]
        if not sid:
            continue
        mine = []
        for col, owner, t, h in columns:
            want = deviation_scores(t.scores).get(sid) if (t and owner == dept) else None
            if want is not None:
                mine.append(want)
            report.add(check)
            if not _same(line[col], want):
                report.findings.append(Finding(check, f"{ws.title} {sid} {h.replace(chr(10), ' ')}",
                                               f"表は{line[col]!r}、計算では{want!r}"))
        report.add(check)
        want = sum(mine) / len(mine) if mine else None
        if not _same(line[-1], want):
            report.findings.append(Finding(check, f"{ws.title} {sid} 平均偏差値", f"表は{line[-1]!r}、計算では{want!r}"))


def verify(output: Path, values_wb, depts: Sequence[DeptSpec], rules: Rules) -> Report:
    """output は式のままのブック、values_wb は再計算した値のブック。"""
    report = Report()
    formulas = load_workbook(output)
    check_errors(values_wb, report)
    for d in depts:
        original = load_workbook(d.original).worksheets[0]
        original_values = load_workbook(d.original, data_only=True).worksheets[0]
        check_original(formulas[f"原本_{d.short}"], original, report)
        for ws in load_workbook(d.register).worksheets:
            if re.search(r"\d+月$", ws.title.strip()) and "原紙" not in ws.title:
                check_original(formulas[ws.title], ws, report, "出席簿シートが元の出席簿と同じか")
        check_untouched(formulas[f"AI計算版_{d.short}"], original, rules, report, rules.corrections.get(d.name))
        check_red(formulas[f"AI計算版_{d.short}"], values_wb[f"AI計算版_{d.short}"], original_values, report)
        check_ai_values(values_wb[f"AI計算版_{d.short}"], original_values, d, rules, report)
    check_linked(values_wb, depts, report)
    return report
