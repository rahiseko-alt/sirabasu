from openpyxl import Workbook

from grading.analysis.sheets import add_deviation_sheet, add_test_analysis, collect_tests


def _grades():
    ws = Workbook().active
    ws["B3"], ws["E3"], ws["E4"], ws["F4"], ws["E5"], ws["F5"] = "学籍番号", "マーケティング", "出席", "筆記テスト", 10, 80
    for i, (sid, score) in enumerate([("A1", 40), ("A2", 48), ("A3", 56), ("A4", 64), ("A5", None)]):
        ws.cell(6 + i, 2, sid)
        ws.cell(6 + i, 3, sid.lower())
        ws.cell(6 + i, 5, 10)
        ws.cell(6 + i, 6, score)
    return ws


def test_only_test_items_are_analysed_and_blank_students_skipped():
    cols = collect_tests(_grades(), "国際ビジネス科")
    assert [(c.subject, c.item, c.maximum, len(c.scores)) for c in cols] == [("マーケティング", "筆記テスト", 80, 4)]


def test_analysis_and_deviation_sheets_have_one_row_per_test_and_student():
    wb = Workbook()
    tests = collect_tests(_grades(), "国際ビジネス科")
    a = add_test_analysis(wb, tests)
    assert a["A4"].value == "国際ビジネス科" and a["D4"].value == "適正" and a["F4"].value == 4 and a["I4"].value == 0.65
    d = add_deviation_sheet(wb, tests)
    assert d.max_row == 6 and d["E2"].value < 50 < d["E5"].value and d["E6"].value is None
