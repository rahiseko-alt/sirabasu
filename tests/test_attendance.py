import datetime as dt
from fractions import Fraction

import pytest
from openpyxl import Workbook

from grading.importing.attendance import read_register, subject_counts, calendar_anomalies


def _month(ws, title, days, students, extra=None):
    """days: [(見出し, [(科目, 時限), ...]), ...]。見出しは文字列か datetime か '=E3+1' 等の式。"""
    ws.title = title
    for col, v in zip("ABCD", ["NO.", "学籍番号", "氏　名", "日本名"]):
        ws[f"{col}3"] = v
    c = 5
    for header, periods in days:
        ws.cell(3, c, header)
        if isinstance(header, (dt.datetime, str)) and not str(header).startswith("　"):
            pass
        for subj, period in periods:
            ws.cell(4, c, subj)
            ws.cell(5, c, period)
            c += 1
    ws.cell(3, c, "授業合計数")
    for i, (sid, rom, kana) in enumerate(students):
        r = 6 + i
        ws.cell(r, 1, i + 1)
        ws.cell(r, 2, sid)
        ws.cell(r, 3, rom)
        ws.cell(r, 4, kana)
    for (r, col), v in (extra or {}).items():
        ws.cell(r, col, v)
    return ws


STUDENTS = [("AIBC26001", "TARO", "タロウ"), ("AIBC26002", "HANA", "ハナ")]


@pytest.fixture
def register(tmp_path):
    wb = Workbook()
    _month(wb.active, "国際ビジネスAI科 4月", [
        ("　4月　　13日（月）", [("ビジネス日本語", 1), ("ビジネス日本語", 2), ("⽇本語運⽤⼒強化演習", 3)]),
        ("　4月　　26日（火）", [("マーケティング ", 1)]),   # 4/26 は日曜：記入ミスの例
        ("　月　　日（水）", [("マーケティング", 1)]),       # 日付の入っていない列
    ], STUDENTS, {(6, 5): "欠", (7, 6): "遅", (6, 7): "×", (6, 8): "欠", (6, 9): "×", (7, 9): "×"})
    ws = wb.create_sheet()
    _month(ws, "国際ビジネスAI科 8月", [
        (dt.datetime(2026, 8, 10), [("ビジネス日本語", 1)]),
        ('=TEXT(E3,"( aaa )")', []),
        ("=E3+1", [("ビジネス日本語", 1)]),   # 8/11 山の日
        ("=E3+2", [("ビジネス日本語", 1)]),   # 8/12（全員×）
    ], STUDENTS, {(6, 6): "欠", (6, 7): "×", (7, 7): "×"})
    path = tmp_path / "reg.xlsx"
    wb.save(path)
    return path


def test_dates_are_read_from_text_datetime_and_simple_formulas(register):
    reg = read_register(register)
    dates = [s.date for s in reg.sessions]
    assert dates[:3] == [dt.date(2026, 4, 13)] * 3
    assert dates[-3:] == [dt.date(2026, 8, 10), dt.date(2026, 8, 11), dt.date(2026, 8, 12)]


def test_subject_names_are_unified_but_the_original_text_is_kept(register):
    reg = read_register(register)
    s = reg.sessions[2]
    assert s.subject == "日本語運用力強化演習" and s.subject_original == "⽇本語運⽤⼒強化演習"
    assert reg.sessions[3].subject == "マーケティング" and reg.sessions[3].subject_original == "マーケティング "


def test_counts_exclude_cross_and_keep_the_cell_of_every_absence_and_late(register):
    counts = subject_counts(read_register(register))
    c = counts[("AIBC26001", "ビジネス日本語")]
    assert (c.sessions, c.absent, c.late) == (4, 2, 0)
    assert c.absent_cells == ("国際ビジネスAI科 4月!E6", "国際ビジネスAI科 8月!F6")
    assert counts[("AIBC26002", "ビジネス日本語")].late_cells == ("国際ビジネスAI科 4月!F7",)
    assert ("AIBC26001", "日本語運用力強化演習") not in counts   # ×しか無い


def test_attendance_rate_follows_the_registers_own_rule_exactly(register):
    c = subject_counts(read_register(register))[("AIBC26002", "ビジネス日本語")]
    assert c.rate == 1 - Fraction(1, 3) / 4       # 遅刻1回＝欠席1/3、丸めない


def test_problems_in_the_headers_are_reported_with_their_cell(register):
    reg = read_register(register)
    texts = [(p.cell, p.detail) for p in reg.problems]
    assert any(cell == "国際ビジネスAI科 4月!H3" and "日曜" in d for cell, d in texts)
    assert any(cell == "国際ビジネスAI科 4月!I3" and "日付" in d for cell, d in texts)


def test_a_dated_correction_replaces_the_header_date(register):
    reg = read_register(register, date_corrections={"国際ビジネスAI科 4月!H3": dt.date(2026, 4, 28)})
    assert reg.sessions[3].date == dt.date(2026, 4, 28)
    assert not any(p.cell == "国際ビジネスAI科 4月!H3" for p in reg.problems)


def test_dateless_column_with_marks_is_a_problem(tmp_path):
    wb = Workbook()
    _month(wb.active, "総合ビジネス科 5月", [("　月　　日（月）", [("AI演習", 1)])], STUDENTS, {(6, 5): "欠"})
    wb.save(tmp_path / "r.xlsx")
    reg = read_register(tmp_path / "r.xlsx")
    assert any("日付が無い" in p.detail for p in reg.problems)


def test_unknown_mark_is_a_problem(tmp_path):
    wb = Workbook()
    _month(wb.active, "総合ビジネス科 4月", [("　4月　　13日（月）", [("AI演習", 1)])], STUDENTS, {(6, 5): "公"})
    wb.save(tmp_path / "r.xlsx")
    assert any("記号" in p.detail for p in read_register(tmp_path / "r.xlsx").problems)


def test_calendar_anomalies_find_holiday_classes_and_weekdays_without_class(register):
    found = calendar_anomalies(read_register(register), first=dt.date(2026, 8, 10), last=dt.date(2026, 8, 14))
    kinds = {(a.date, a.kind) for a in found}
    assert (dt.date(2026, 8, 11), "祝日に授業の記録") in kinds
    assert (dt.date(2026, 8, 12), "平日なのに全員×") in kinds
    assert (dt.date(2026, 8, 13), "平日なのに出席簿に日付が無い") in kinds


def test_weekday_with_a_date_but_no_subjects_is_reported_as_such(tmp_path):
    wb = Workbook()
    _month(wb.active, "国際ビジネスAI科 8月", [(dt.datetime(2026, 8, 4), [(None, 1)])], STUDENTS)
    wb.save(tmp_path / "r.xlsx")
    found = calendar_anomalies(read_register(tmp_path / "r.xlsx"), dt.date(2026, 8, 4), dt.date(2026, 8, 4))
    assert [a.kind for a in found] == ["平日なのに科目の記入が無い"]


def test_subject_summary_keeps_layout_counts_and_shows_differences(register, tmp_path):
    from openpyxl import load_workbook
    from grading.export.attendance_sheet import write_subject_summary

    wb = load_workbook(register)
    ws = wb.create_sheet("科目ごと")
    ws["A1"] = "2026年度　国際ビジネスAI科"
    ws["E3"] = "ビジネス日本語"
    for k, h in enumerate(["授業合計数", "欠席合計数", "遅刻合計数", "欠席数", "出席率"]):
        ws.cell(4, 5 + k, h)
    for i, (sid, rom, kana) in enumerate(STUDENTS):
        ws.cell(6 + i, 2, sid)
        ws.cell(6 + i, 5, "=999")
    wb.save(register)
    counts = subject_counts(read_register(register))
    out = write_subject_summary(register, tmp_path / "o.xlsx", counts,
                                original={("AIBC26001", "ビジネス日本語"): (5, 2, 0)})
    wb = load_workbook(out)
    s = wb.worksheets[0]
    assert wb.sheetnames == [s.title, "元の集計との比較", "根拠", "日程と記入の確認"]
    assert [s.cell(6, c).value for c in range(5, 10)] == [4, 2, 0, "=F6+G6/3", "=1-H6/E6"]
    assert s["A1"].value.endswith("（AI計算版）")
    diff = [r for r in wb["元の集計との比較"].iter_rows(min_row=2, values_only=True) if r[0] == "AIBC26001"]
    assert diff[0][2:4] == (5, 4) and diff[0][8] == "違う"
    assert ("AIBC26001", "ビジネス日本語", "欠", "国際ビジネスAI科 4月!E6") in list(wb["根拠"].iter_rows(values_only=True))


def _weekly_register(tmp_path, marks):
    wb = Workbook()
    _month(wb.active, "国際ビジネスAI科 5月", [
        ("　5月　　11日（月）", [("⽇本語運⽤⼒強化演習", 5)]),
        ("　5月　　12日（火）", [("⽇本語運⽤⼒強化演習", 5)]),
        ("　5月　　18日（月）", [("⽇本語運⽤⼒強化演習", 5)]),
        ("　5月　　19日（火）", [("⽇本語運⽤⼒強化演習", 5)]),
    ], STUDENTS, marks)
    wb.save(tmp_path / "w.xlsx")
    return read_register(tmp_path / "w.xlsx")


def test_weekly_subject_counts_a_week_once_if_any_session_was_attended(tmp_path):
    from grading.importing.attendance import weekly_counts
    # 学生1: 1週目 月欠・火出 → 出席1回、2週目 月×・火欠 → 欠席1回
    # 学生2: 1週目 月出・火× → 出席、2週目 月遅・火× → 遅刻
    reg = _weekly_register(tmp_path, {(6, 5): "欠", (6, 7): "×", (6, 8): "欠", (7, 6): "×", (7, 7): "遅", (7, 8): "×"})
    c = weekly_counts(reg, "日本語運用力強化演習")
    one, two = c[("AIBC26001", "日本語運用力強化演習")], c[("AIBC26002", "日本語運用力強化演習")]
    assert (one.sessions, one.absent, one.late) == (2, 1, 0)
    assert (two.sessions, two.absent, two.late) == (2, 0, 1)
    assert one.absent_cells == ("国際ビジネスAI科 5月!H6",)


def test_week_with_only_cross_is_not_counted(tmp_path):
    from grading.importing.attendance import weekly_counts
    reg = _weekly_register(tmp_path, {(6, 5): "×", (6, 6): "×"})
    assert weekly_counts(reg, "日本語運用力強化演習")[("AIBC26001", "日本語運用力強化演習")].sessions == 1


def test_weekly_counts_can_leave_out_named_sessions(tmp_path):
    from grading.importing.attendance import weekly_counts
    reg = _weekly_register(tmp_path, {(6, 5): "欠", (6, 6): "×"})
    full = weekly_counts(reg, "日本語運用力強化演習")[("AIBC26001", "日本語運用力強化演習")]
    cut = weekly_counts(reg, "日本語運用力強化演習", exclude={"国際ビジネスAI科 5月!E"})[("AIBC26001", "日本語運用力強化演習")]
    assert (full.sessions, full.absent) == (2, 1) and (cut.sessions, cut.absent) == (1, 0)
