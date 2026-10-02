from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from grading.export.copy_sheet import copy_sheet


def test_sheet_is_copied_with_values_formulas_styles_merges_and_widths():
    src = Workbook().active
    src["A1"] = "2026年度　総合ビジネス科"
    src.merge_cells("A1:D1")
    src["B6"], src["E6"], src["F6"] = "AIBC26003", 8, "=SUM(E6:E6)"
    src["E6"].font = Font(bold=True)
    src["E6"].fill = PatternFill("solid", fgColor="FFFF00")
    src["E6"].number_format = "0.0"
    src.column_dimensions["B"].width = 17
    src.row_dimensions[6].height = 30
    src.freeze_panes = "E6"
    dst_wb = Workbook()
    out = copy_sheet(src, dst_wb, "原本_総合")
    assert out.title == "原本_総合"
    assert (out["A1"].value, out["B6"].value, out["E6"].value, out["F6"].value) == ("2026年度　総合ビジネス科", "AIBC26003", 8, "=SUM(E6:E6)")
    assert out["E6"].font.bold and out["E6"].fill.fgColor.rgb.endswith("FFFF00") and out["E6"].number_format == "0.0"
    assert [str(r) for r in out.merged_cells.ranges] == ["A1:D1"]
    assert out.column_dimensions["B"].width == 17 and out.row_dimensions[6].height == 30 and out.freeze_panes == "E6"
