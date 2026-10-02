from openpyxl import Workbook

from grading.analysis.distribution import _bins, add_distribution_sheet, student_totals
from grading.export.e_list import GradeRow


def test_bins_put_the_top_score_in_the_last_band():
    labels, counts = _bins([0, 49, 50, 99, 100], 50, 100)
    assert labels == ["0〜49", "50〜100"] and counts == [2, 3]


def test_student_totals_sum_subject_totals():
    rows = [GradeRow("A", "S1", "x", 60, "D"), GradeRow("B", "S1", "x", 70.5, "C"), GradeRow("A", "S2", "y", None, None)]
    assert student_totals(rows) == {"S1": 130.5}


def test_distribution_sheet_has_tables_and_charts():
    rows = [GradeRow(s, sid, sid, t, "C") for sid, ts in (("S1", (60, 70)), ("S2", (80, 95))) for s, t in zip("AB", ts)]
    wb = Workbook()
    ws = add_distribution_sheet(wb, [("国際ビジネス科", rows)])
    assert ws["A1"].value.startswith("総点の度数分布（2科目・200点満点）")
    assert [ws.cell(r, 1).value for r in range(3, 7)] == ["0〜49", "50〜99", "100〜149", "150〜200"]
    assert [ws.cell(r, 2).value for r in range(3, 7)] == [0, 0, 1, 1]
    assert len(ws._charts) == 2
