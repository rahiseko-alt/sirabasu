from openpyxl import Workbook, load_workbook

from grading.export.e_list import add_e_matrix, grade_rows


def _sheet():
    wb = Workbook()
    ws = wb.active
    ws["B3"], ws["E3"], ws["I3"] = "学籍番号", "マーケティング", "総合ビジネス概論"
    for c, h in zip("EFGH", ["出席", "授業態度", "合計", "評定"]):
        ws[f"{c}4"] = h
    for c, h in zip("IJK", ["学習参加", "合計", "評定"]):
        ws[f"{c}4"] = h
    ws["B6"], ws["C6"], ws["E6"], ws["F6"], ws["G6"], ws["H6"], ws["J6"], ws["K6"] = "AIBC26001", "TARO", 5, 4, 9, "E", 0, "E"
    ws["B7"], ws["C7"], ws["E7"], ws["F7"], ws["G7"], ws["H7"], ws["J7"], ws["K7"] = "AIBC26002", "HANA", 50, 40, 90, "A", 0, "E"
    ws["I7"] = 3
    return ws


def test_e_rows_skip_subjects_the_student_does_not_take():
    rows = grade_rows(_sheet())
    es = [(r.subject, r.student_id) for r in rows if r.grade == "E"]
    assert es == [("マーケティング", "AIBC26001"), ("総合ビジネス概論", "AIBC26002")]


def test_e_list_per_department_is_a_student_by_subject_table(tmp_path):
    ws = _sheet()
    ws["D6"], ws["D7"] = "タロウ\u3000ヤマ\nダ", "ハナ"
    wb = Workbook()
    info = {"マーケティング": ("百井", "月"), "総合ビジネス概論": ("元島", "金")}
    add_e_matrix(wb, [("国際ビジネス科", grade_rows(ws), info)], "E一覧_国際")
    wb.save(tmp_path / "e.xlsx")
    m = load_workbook(tmp_path / "e.xlsx")["E一覧_国際"]
    assert [c.value for c in m[1]][4:] == ["マーケティング", "総合ビジネス概論"]
    assert [c.value for c in m[2]][4:] == ["百井", "元島"] and [c.value for c in m[3]][4:] == ["月", "金"]
    assert [c.value for c in m[4]][3:] == [2, 1, 1]
    assert [c.value for c in m[6]] == ["AIBC26001", "TARO", "タロウ ヤマ ダ", 1, 9, None]
    assert [c.value for c in m[7]] == ["AIBC26002", "HANA", "ハナ", 1, None, 0]


def test_combined_e_list_adds_the_department_and_marks_differing_teachers(tmp_path):
    a, b = _sheet(), _sheet()
    b["B6"], b["B7"] = "AIBC26003", "AIBC26004"
    wb = Workbook()
    add_e_matrix(wb, [("国際ビジネス科", grade_rows(a), {"マーケティング": ("百井", "月")}),
                      ("総合ビジネス科", grade_rows(b), {"マーケティング": ("小齊平", "月")})], "E一覧_統合")
    wb.save(tmp_path / "e.xlsx")
    m = load_workbook(tmp_path / "e.xlsx")["E一覧_統合"]
    assert m["A5"].value == "学科" and m["F2"].value == "百井（国際）／小齊平（総合）" and m["F3"].value == "月"
    assert [m.cell(r, 1).value for r in range(6, 10)] == ["国際ビジネス科", "国際ビジネス科", "総合ビジネス科", "総合ビジネス科"]
    assert m["E4"].value == 4


def test_personal_grade_table_has_one_row_per_student(tmp_path):
    from grading.export.e_list import add_personal_grades
    ws = _sheet()
    ws["L3"], ws["L6"], ws["L7"] = "GPA", 1.0, 3.5
    wb = Workbook()
    add_personal_grades(wb, grade_rows(ws), gpa_of(ws))
    wb.save(tmp_path / "p.xlsx")
    out = load_workbook(tmp_path / "p.xlsx")["個人別評定"]
    assert [c.value for c in out[1]] == ["学籍番号", "氏名", "カタカナ", "マーケティング", "総合ビジネス概論", "GPA", "GPA評定"]
    assert [c.value for c in out[2]] == ["AIBC26001", "TARO", None, "E", "—", 1.0, None]
    assert [c.value for c in out[3]] == ["AIBC26002", "HANA", None, "A", "E", 3.5, None]


from grading.export.e_list import gpa_of  # noqa: E402
