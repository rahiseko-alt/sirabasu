from fractions import Fraction

import pytest
from openpyxl import Workbook, load_workbook

from grading.export.fill import FilledValue, fill_template


def _blank(path):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "2026年度　国際ビジネス科（AI計算版）"
    ws["B3"] = "学籍番号"
    ws["E3"], ws["E4"], ws["F4"], ws["G4"] = "マーケティング", "出席", "授業態度", "筆記テスト"
    ws["H3"], ws["H4"], ws["I4"] = "AI演習（実践）", "出席", "授業態度"
    ws.merge_cells("E3:G3")
    ws["B6"], ws["B7"] = "AIBC26001", "AIBC26002"
    wb.save(path)
    return path


def _v(sid, subject, item, value, ev=("出席簿!F6",)):
    return FilledValue(sid, subject, item, Fraction(value), "10 − 2 × 遅刻1回 = 8", tuple(ev))


def test_values_go_to_the_cell_found_by_subject_and_item(tmp_path):
    out = fill_template(_blank(tmp_path / "b.xlsx"), tmp_path / "o.xlsx",
                        [_v("AIBC26001", "マーケティング", "授業態度", 8), _v("AIBC26002", "AI演習", "授業態度", 10, ())])
    ws = load_workbook(out).worksheets[0]
    assert ws["F6"].value == 8 and ws["I7"].value == 10
    assert ws["E6"].value is None


def test_every_value_is_listed_with_its_formula_and_source(tmp_path):
    out = fill_template(_blank(tmp_path / "b.xlsx"), tmp_path / "o.xlsx", [_v("AIBC26001", "マーケティング", "授業態度", 8)])
    rows = list(load_workbook(out)["根拠"].iter_rows(min_row=2, values_only=True))
    assert rows == [("AIBC26001", "マーケティング", "授業態度", "F6", 8, "8", "10 − 2 × 遅刻1回 = 8", "出席簿!F6")]


def test_fractions_are_kept_exactly_in_the_evidence_sheet(tmp_path):
    out = fill_template(_blank(tmp_path / "b.xlsx"), tmp_path / "o.xlsx", [_v("AIBC26001", "マーケティング", "筆記テスト", Fraction(1, 3))])
    row = next(load_workbook(out)["根拠"].iter_rows(min_row=2, values_only=True))
    assert row[5] == "1/3"


@pytest.mark.parametrize("sid, subject, item", [("AIBC26999", "マーケティング", "授業態度"),
                                                ("AIBC26001", "経営学", "授業態度"),
                                                ("AIBC26001", "マーケティング", "発表")])
def test_unknown_student_subject_or_item_stops_instead_of_guessing(tmp_path, sid, subject, item):
    with pytest.raises(KeyError):
        fill_template(_blank(tmp_path / "b.xlsx"), tmp_path / "o.xlsx", [_v(sid, subject, item, 8)])


def test_writing_the_same_cell_twice_is_refused(tmp_path):
    with pytest.raises(ValueError):
        fill_template(_blank(tmp_path / "b.xlsx"), tmp_path / "o.xlsx",
                      [_v("AIBC26001", "マーケティング", "授業態度", 8), _v("AIBC26001", "マーケティング", "授業態度", 6)])


def _original(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "２０２６年前期"
    ws["A1"] = "2026年度　国際ビジネス科"
    ws["B3"] = "学籍番号"
    ws["E3"], ws["E4"], ws["F4"], ws["G4"], ws["H4"] = "マーケティング", "出席", "授業態度", "筆記テスト", "合計"
    ws["I3"], ws["I4"], ws["J4"] = "ビジネス日本語", "出席", "授業態度"
    ws["B6"] = "AIBC26001"
    ws["E6"], ws["F6"], ws["G6"], ws["H6"] = "=10-((1-0.9)/0.04)", 8, 56, "=SUM(E6:G6)"
    ws["I6"], ws["J6"] = 7, 18
    wb.save(path)
    return path


def test_only_ai_cells_are_replaced_and_coloured(tmp_path):
    from grading.export.fill import apply_ai_values
    ws = load_workbook(_original(tmp_path / "o.xlsx")).worksheets[0]
    n = apply_ai_values(ws, [_v("AIBC26001", "マーケティング", "出席", 9), _v("AIBC26001", "ビジネス日本語", "出席", 10)])
    assert n == 2
    assert (ws["E6"].value, ws["I6"].value) == (9, 10)
    assert (ws["F6"].value, ws["G6"].value, ws["H6"].value, ws["J6"].value) == (8, 56, "=SUM(E6:G6)", 18)
    assert ws["E6"].fill.fgColor.rgb.endswith("DDEBF7") and not ws["F6"].fill.fgColor.rgb.endswith("DDEBF7")


def test_unknown_item_stops(tmp_path):
    from grading.export.fill import apply_ai_values
    ws = load_workbook(_original(tmp_path / "o.xlsx")).worksheets[0]
    with pytest.raises(KeyError):
        apply_ai_values(ws, [_v("AIBC26001", "マーケティング", "発表", 9)])


def test_rate_columns_come_from_the_original_formula_or_the_same_position(tmp_path):
    from grading.export.fill import apply_rates, rate_columns
    ws = load_workbook(_original(tmp_path / "o.xlsx")).worksheets[0]
    ws["M6"] = 0.9     # ビジネス日本語の出席(I)＋4 の位置に率がある（式なし）
    cols = rate_columns(ws)
    assert cols == {"ビジネス日本語": 13}   # マーケの＋4の位置（I列）は7で率ではないので使わない
    ws["E6"] = "=10-((1-K6)/0.04)"
    assert rate_columns(ws)["マーケティング"] == 11
    n = apply_rates(ws, rate_columns(ws), {("AIBC26001", "マーケティング"): Fraction(17, 18)})
    assert n == 1 and ws["K6"].value == float(Fraction(17, 18)) and ws["K6"].fill.fgColor.rgb.endswith("DDEBF7")


def _totals_sheet():
    from openpyxl import Workbook
    ws = Workbook().active
    ws["E3"], ws["E4"], ws["F4"], ws["G4"], ws["H4"] = "ビジネス演習（理論）", "出席", "テスト", "合計", "評定"
    ws["B6"], ws["E6"], ws["F6"], ws["G6"] = "AIBC26001", 20, 40, "=SUM(E6:I6)"
    ws["B7"], ws["E7"], ws["F7"], ws["G7"] = "AIBC26002", 20, 40, "=SUM(E7:F7)"
    return ws


def test_rate_column_is_taken_out_of_the_total_and_marked():
    from grading.export.fill import FIX_FILL, exclude_rates_from_totals
    ws = _totals_sheet()
    fixed = exclude_rates_from_totals(ws, {"ビジネス演習(理論)": 9})
    assert (ws["G6"].value, ws["G7"].value) == ("=SUM(E6:H6)", "=SUM(E7:F7)")
    assert fixed == [("G6", "=SUM(E6:I6)", "=SUM(E6:H6)")]
    assert ws["G6"].fill.fgColor.rgb.endswith(FIX_FILL.fgColor.rgb[-6:])


def test_rate_column_in_the_middle_of_a_total_stops():
    from grading.export.fill import exclude_rates_from_totals
    ws = _totals_sheet()
    ws["G6"] = "=SUM(E6:K6)"
    with pytest.raises(ValueError):
        exclude_rates_from_totals(ws, {"ビジネス演習(理論)": 9})


def test_cells_that_differ_from_the_original_turn_red_with_white_text():
    from openpyxl import Workbook
    from grading.export.fill import mark_differences
    ai, ai_values, original = Workbook().active, Workbook().active, Workbook().active
    for ws in (ai_values, original):
        ws["B6"], ws["E6"], ws["F6"] = "AIBC26001", 8, 50
    ai_values["E6"], ai["E6"], ai["F6"] = 9, "=X", "=Y"
    assert mark_differences(ai, ai_values, original) == ["E6"]
    assert ai["E6"].fill.fgColor.rgb.endswith("FF0000") and ai["E6"].font.color.rgb.endswith("FFFFFF")
    assert not ai["F6"].fill.fgColor.rgb.endswith("FF0000")


def test_a_correction_is_written_to_the_cell_found_by_subject_and_item():
    from openpyxl import Workbook
    from grading.export.fill import apply_corrections
    ws = Workbook().active
    ws["E3"], ws["E4"], ws["F4"] = "ビジネス演習（実践）", "実技テスト", "筆記課題\n筆記テスト"
    ws["B6"], ws["F6"] = "AIBC26018", 28
    assert apply_corrections(ws, {("AIBC26018", "ビジネス演習(実践)", "筆記課題\n筆記テスト"): 25}) == ["F6"]
    assert ws["F6"].value == 25
    with pytest.raises(KeyError):
        apply_corrections(ws, {("AIBC26099", "ビジネス演習(実践)", "筆記課題\n筆記テスト"): 25})
