import datetime as dt
from fractions import Fraction

import pytest
from openpyxl import Workbook

from grading.export.report_card import ReportCardError, read_cards, render_html
from grading.importing.attendance import SubjectCount


def _sheet():
    """AI計算版と同じ並び。2人、科目はビジネス日本語と国際理解と総合ビジネス概論（誰も受けていない）。"""
    ws = Workbook().active
    ws["B3"], ws["E3"], ws["H3"], ws["K3"] = "学籍番号", "ビジネス日本語", "国際理解", "総合ビジネス概論"
    for cols in ("EFG", "HIJ", "KLM"):
        for c, h in zip(cols, ["出席", "合計", "評定"]):
            ws[f"{c}4"] = h
    ws["B6"], ws["C6"], ws["D6"] = "AIBC26001", "TARO YAMADA", "タロウ　ヤマ\nダ"
    ws["E6"], ws["F6"], ws["G6"], ws["H6"], ws["I6"], ws["J6"] = 8, 280 / 9, "E", 9, 85, "A"
    ws["B7"], ws["C7"], ws["D7"] = "AIBC26002", "HANA SATO", "ハナ"
    ws["E7"], ws["F7"], ws["G7"], ws["H7"], ws["I7"], ws["J7"] = 10, 90, "A", 7, 70.05, "B"
    return ws


def _counts():
    c = {}
    for sid, (abs1, late1) in {"AIBC26001": (3, 2), "AIBC26002": (0, 0)}.items():
        c[(sid, "ビジネス日本語")] = SubjectCount(sid, "ビジネス日本語", 30, abs1, late1, (), ())
        c[(sid, "国際社会I:異文化理解")] = SubjectCount(sid, "国際社会I:異文化理解", 20, 0, 0, (), ())
    return c


def test_card_lists_only_subjects_the_student_takes_in_template_order_with_credits():
    card = read_cards(_sheet(), "国際ビジネス科", _counts())[0]
    assert (card.student_id, card.name_en, card.name_ja) == ("AIBC26001", "TARO YAMADA", "タロウ ヤマ ダ")
    assert [(l.name, l.grade, l.credits, l.earned) for l in card.lines] == [
        ("ビジネス日本語Ⅰ", "E", 2, None), ("国際社会・異文化理解Ⅰ", "A", 2, 2)]
    assert (card.credits_set, card.credits_earned) == (4, 2)


def test_attendance_counts_three_lates_as_one_absence_without_rounding():
    card = read_cards(_sheet(), "国際ビジネス科", _counts())[0]
    assert card.hours == 50
    assert card.attended == 50 - 3 - Fraction(2, 3)


def test_missing_attendance_stops_instead_of_guessing():
    counts = _counts()
    del counts[("AIBC26002", "国際社会I:異文化理解")]
    with pytest.raises(ReportCardError, match="AIBC26002"):
        read_cards(_sheet(), "国際ビジネス科", counts)


def test_unknown_subject_stops():
    ws = _sheet()
    ws["H3"] = "未知の科目"
    with pytest.raises(ReportCardError, match="未知の科目"):
        read_cards(ws, "国際ビジネス科", _counts())


def test_grade_outside_abcde_stops():
    ws = _sheet()
    ws["J7"] = None
    with pytest.raises(ReportCardError, match="AIBC26002"):
        read_cards(ws, "国際ビジネス科", _counts())


def test_html_shows_one_decimal_scores_issue_date_and_dash_for_failed_credit():
    cards = read_cards(_sheet(), "国際ビジネス科", _counts())
    page = render_html(cards, dt.date(2026, 10, 2))
    assert page.count('class="page"') == 2
    assert "31.1" in page and "70.1" in page and ">90<" in page
    assert "2026年10月2日" in page and "－" in page
    assert "この書類は成績証明書ではありません" in page
    assert "46.3" in page          # 出席時数 50−3−2/3 = 46.33…
    assert "92.7%" in page         # 46.33… ÷ 50


def test_e_rows_can_be_black_with_white_text_and_others_stay_plain():
    cards = read_cards(_sheet(), "国際ビジネス科", _counts())
    page = render_html(cards[:1], dt.date(2026, 10, 2), black_e=True)
    assert page.count('<tr class="fail">') == 1        # ビジネス日本語のE だけ
    assert render_html(cards[:1], dt.date(2026, 10, 2)).count('class="fail"') == 0


def test_stamp_goes_only_on_cards_with_an_e():
    cards = read_cards(_sheet(), "国際ビジネス科", _counts())
    page = render_html(cards, dt.date(2026, 10, 2), black_e=True, stamp="単位不認定")
    assert page.count('<div class="stamp body">単位不認定</div>') == 1     # E のある1人目だけ
    page = render_html(cards, dt.date(2026, 10, 2), black_e=True, stamp="あなたは不合格です", place="title")
    assert page.count('<div class="stamp title long">あなたは不合格です</div>') == 1


def test_solid_stamp_is_black_with_white_text():
    cards = read_cards(_sheet(), "国際ビジネス科", _counts())
    page = render_html(cards, dt.date(2026, 10, 2), black_e=True, stamp="不合格", place="title", solid=True)
    assert page.count('<div class="stamp title solid">不合格</div>') == 1


def test_certificate_has_its_title_and_no_not_a_certificate_notice():
    cards = read_cards(_sheet(), "国際ビジネス科", _counts())
    page = render_html(cards, dt.date(2026, 10, 2), kind="成績証明書")
    assert "<h1>成績証明書</h1>" in page and "<h1>通知表</h1>" not in page
    assert "この書類は成績証明書ではありません" not in page
    assert 'class="fail"' not in page and 'class="stamp' not in page
