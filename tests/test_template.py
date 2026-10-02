from openpyxl import Workbook, load_workbook

from grading.export.template import blank_template


def _school_sheet(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "２０２６年前期"
    ws["A1"] = "2026年度　国際ビジネス科"
    for col, v in zip("ABCD", ["NO.", "学籍番号", "氏　名", "科目名"]):
        ws[f"{col}3"] = v
    ws["E3"] = "マーケティング"
    for col, v in zip("EFGHIJK", ["出席", "授業態度", "筆記テスト", None, None, "合計", "評定"]):
        ws[f"{col}4"] = v
    ws["D4"], ws["D5"] = "評価項目", "点数配分"
    ws["E5"], ws["F5"], ws["G5"], ws["H5"], ws["J5"] = 10, 10, 80, 0.8, "=SUM(E5:I5)"
    ws.append([])  # row 6 below
    ws["A6"], ws["B6"], ws["C6"], ws["D6"] = 1, "AIBC26001", "TARO", "タロウ"
    ws["E6"], ws["F6"], ws["G6"], ws["H6"], ws["I6"] = "=10-((1-I6)/0.04)", 8, "=ROUNDDOWN(H6*0.8,0)", 70, 0.9
    ws["J6"], ws["K6"] = "=SUM(E6:G6)", '=LOOKUP(J6,{0,60},{"E","D"})'
    ws["A7"], ws["B7"], ws["C7"], ws["D7"] = 2, "AIBC26002", "HANAKO", "ハナコ"
    ws["F7"] = 10
    ws["E6"].number_format = "0.0"
    wb.save(path)
    return path


def test_everything_from_column_e_in_student_rows_is_blanked(tmp_path):
    out = blank_template(_school_sheet(tmp_path / "src.xlsx"), tmp_path / "out.xlsx")
    ws = load_workbook(out).active
    assert all(ws.cell(r, c).value is None for r in (6, 7) for c in range(5, 12))


def test_student_identity_and_headings_are_kept(tmp_path):
    ws = load_workbook(blank_template(_school_sheet(tmp_path / "s.xlsx"), tmp_path / "o.xlsx")).active
    assert [ws.cell(6, c).value for c in range(1, 5)] == [1, "AIBC26001", "TARO", "タロウ"]
    assert ws["E3"].value == "マーケティング" and ws["F4"].value == "授業態度" and ws["D5"].value == "点数配分"


def test_allocations_under_named_items_are_kept_but_formulas_and_helpers_are_removed(tmp_path):
    ws = load_workbook(blank_template(_school_sheet(tmp_path / "s.xlsx"), tmp_path / "o.xlsx")).active
    assert (ws["E5"].value, ws["F5"].value, ws["G5"].value) == (10, 10, 80)
    assert ws["H5"].value is None   # 見出しの無い補助列の係数
    assert ws["J5"].value is None   # 合計の式


def test_no_formula_remains_anywhere(tmp_path):
    ws = load_workbook(blank_template(_school_sheet(tmp_path / "s.xlsx"), tmp_path / "o.xlsx")).active
    assert not [c.coordinate for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]


def test_title_is_marked_as_the_ai_version_and_source_is_untouched(tmp_path):
    src = _school_sheet(tmp_path / "s.xlsx")
    ws = load_workbook(blank_template(src, tmp_path / "o.xlsx")).active
    assert ws["A1"].value.endswith("（AI計算版）")
    assert load_workbook(src).active["E6"].value == "=10-((1-I6)/0.04)"


def test_formatting_is_kept(tmp_path):
    ws = load_workbook(blank_template(_school_sheet(tmp_path / "s.xlsx"), tmp_path / "o.xlsx")).active
    assert ws["E6"].number_format == "0.0"


def test_merged_heading_cells_are_left_alone(tmp_path):
    src = _school_sheet(tmp_path / "s.xlsx")
    wb = load_workbook(src)
    ws = wb.active
    ws["L3"] = "GPA"
    ws.merge_cells("L3:L5")
    ws["L6"] = "=ROUNDDOWN(1/12,1)"
    wb.save(src)
    ws = load_workbook(blank_template(src, tmp_path / "o.xlsx")).active
    assert ws["L3"].value == "GPA" and ws["L6"].value is None
