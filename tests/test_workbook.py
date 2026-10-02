import shutil
import subprocess

import pytest
from openpyxl import load_workbook

from grading.sample import build

SHEETS = ["説明", "成績表", "検算", "完全性", "根拠", "採点ルール", "確認事項"]


@pytest.fixture(scope="module")
def sample(tmp_path_factory):
    return build(tmp_path_factory.mktemp("wb") / "sample.xlsx")


def test_workbook_has_all_sheets_in_reading_order(sample):
    assert load_workbook(sample).sheetnames == SHEETS


def test_unresolved_homonym_makes_the_workbook_incomplete(sample):
    guide = [c.value for c in load_workbook(sample)["説明"]["A"]]
    assert any("未完成" in (v or "") for v in guide)
    notices = [r[4] for r in load_workbook(sample)["確認事項"].iter_rows(min_row=2, values_only=True)]
    assert any("同姓同名" in (t or "") and "どの学生" in t for t in notices)


def test_every_score_has_a_source_cell(sample):
    rows = list(load_workbook(sample)["根拠"].iter_rows(min_row=2, values_only=True))
    assert rows and all(r[6] and r[8] for r in rows)


def test_grade_that_rounding_could_change_is_listed(sample):
    notices = [r for r in load_workbook(sample)["確認事項"].iter_rows(min_row=2, values_only=True)]
    assert any(r[0] == "要確認" and "端数" in r[4] for r in notices)


def test_totals_are_not_rounded(sample):
    grades = {(r[0], r[3]): r for r in load_workbook(sample)["成績表"].iter_rows(min_row=2, values_only=True)}
    assert grades[("20260082", "AI基礎")][5] == "986/15（65.733333…）"


def _has_calc() -> bool:
    return shutil.which("soffice") is not None and subprocess.run(
        ["dpkg", "-s", "libreoffice-calc"], capture_output=True).returncode == 0


@pytest.mark.skipif(not _has_calc(), reason="LibreOffice Calc が無いため Excel 式の再計算を確認できない")
def test_excel_itself_agrees_with_every_system_value(sample, tmp_path):
    out = tmp_path / "recalc"
    subprocess.run(["soffice", f"-env:UserInstallation=file://{tmp_path}/profile", "--headless", "--convert-to",
                    "xlsx:Calc MS Excel 2007 XML", "--outdir", str(out), str(sample)],
                   check=True, capture_output=True, timeout=180)
    ws = load_workbook(out / sample.name, data_only=True)["検算"]
    verdicts = [ws.cell(r, 15).value for r in range(2, ws.max_row + 1) if ws.cell(r, 15).value]
    assert len(verdicts) >= 20 and set(verdicts) == {"一致"}


def test_html_report_shows_banner_formulas_sources_and_recheck(tmp_path):
    from grading.sample import build_html
    text = build_html(tmp_path / "r.html").read_text(encoding="utf-8")
    assert "未完成" in text
    assert "(8 + 9 + 10) ÷ (10 + 10 + 10) × 40 = <b>36</b>" in text
    assert "AI基礎_課題一覧.xlsx / 課題 / H31" in text
    assert "一致（別の手順で計算し直しても同じ結果）" in text and "不一致" not in text
    assert "値が無い: レポート1、レポート2、発表、参加" in text


def test_html_escapes_text_from_documents(tmp_path):
    from grading.export.report_html import _e
    assert _e("<script>") == "&lt;script&gt;"
