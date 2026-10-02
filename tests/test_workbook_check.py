from openpyxl import Workbook

from grading.validation.workbook_check import (Finding, Report, Rules, check_original, check_untouched,
                                               expected_points, raw_counts)


def test_points_are_full_until_the_rate_drops_by_four_percent():
    assert expected_points(50, 1, 0, 10, 1) == 10      # 98%
    assert expected_points(25, 1, 0, 10, 1) == 9       # 96%
    assert expected_points(25, 10, 0, 10, 1) == 0      # 60%: 4%×10回 → 0
    assert expected_points(10, 4, 1, 10, 1) == 0       # 60%未満
    assert expected_points(25, 1, 0, 20, 2) == 18      # 理論は2点ずつ


def test_three_lates_count_as_one_absence():
    assert expected_points(25, 0, 3, 10, 1) == expected_points(25, 1, 0, 10, 1)


def test_raw_counts_skip_crossed_out_sessions(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "4月"
    ws["E4"], ws["F4"], ws["G4"], ws["H3"] = "マーケティング", "マーケティング", "ＡＩ演習 ", "授業合計数"
    ws["B6"], ws["E6"], ws["F6"], ws["G6"] = "AIBC26001", "欠", "×", "遅"
    wb.save(tmp_path / "r.xlsx")
    counts = raw_counts(tmp_path / "r.xlsx")
    assert counts[("AIBC26001", "マーケティング")] == (1, 1, 0)
    assert counts[("AIBC26001", "AI演習")] == (1, 0, 1)


def _sheet():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "成績表"
    ws["E3"], ws["E4"], ws["F4"], ws["G4"] = "マーケティング", "出席", "授業態度", "筆記テスト"
    ws["B6"], ws["E6"], ws["F6"], ws["G6"] = "AIBC26001", "=10-((1-I6)/0.04)", 9, 50
    return ws


RULES = Rules({"マーケティング": (10, 1)}, "日本語運用力強化演習", ["マーケティング"], [])


def test_only_the_ai_cells_may_differ_from_the_original():
    original, ai = _sheet(), _sheet()
    ai["A1"] = "成績表（AI計算版）"
    ai["E6"], ai["F6"] = "='出席集計_国際'!$G$4", "=10-2*'出席集計_国際'!$F$4"
    ai["I6"] = "='出席集計_国際'!$G$4"
    report = Report()
    check_untouched(ai, original, RULES, report)
    assert report.ok
    ai["G6"] = 51
    check_untouched(ai, original, RULES, report)
    assert [f.where for f in report.findings] == ["Sheet!G6"]


def test_a_changed_formula_in_the_original_copy_is_reported():
    original, copy = _sheet(), _sheet()
    copy["E6"] = 10
    report = Report()
    check_original(copy, original, report)
    assert not report.ok and report.findings[0].where == "Sheet!E6"


def test_report_says_pass_or_fail_in_plain_words():
    report = Report()
    report.add("出席点", 3)
    assert report.text().startswith("検算結果: 合格")
    report.findings.append(Finding("出席点", "AIBC26001", "表は7、数え直すと8"))
    assert "不合格（食い違い1件）" in report.text() and "× 出席点: 3件を照合、食い違い1件" in report.text()


def test_a_pasted_number_in_an_ai_cell_is_reported():
    original, ai = _sheet(), _sheet()
    ai["A1"] = "成績表（AI計算版）"
    ai["E6"], ai["F6"] = 8, "=10-2*'出席集計_国際'!$F$4"
    ai["I6"] = "='出席集計_国際'!$G$4"
    report = Report()
    check_untouched(ai, original, RULES, report)
    assert [(f.check, f.where) for f in report.findings] == [("AIのセルが数字の貼り付けでなく式か", "Sheet!E6")]
